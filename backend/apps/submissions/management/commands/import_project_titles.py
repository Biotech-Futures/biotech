"""Fill in project titles for a year's teams from a spreadsheet of team codes and titles.

    python manage.py import_project_titles titles.xlsx --year 2026 [--dry-run]
"""
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from apps.groups.models import Groups
from apps.submissions.models import Submission

CODE_HEADER = "team code"
TITLE_HEADER = "project title"
MAX_TITLE_LENGTH = Submission._meta.get_field("project_title").max_length


def read_rows(path: str) -> list[tuple[int, str, str]]:
    """(sheet row number, team code, title) for each filled row of the first sheet."""
    try:
        book = load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError) as error:
        raise CommandError(f"Could not open {path}: {error}") from error
    try:
        all_rows = list(book.worksheets[0].iter_rows(values_only=True))
    finally:
        book.close()

    rows = iter(all_rows)
    headers = [str(cell or "").strip().lower() for cell in next(rows, ())]
    if CODE_HEADER not in headers or TITLE_HEADER not in headers:
        raise CommandError('The first row must have "Team code" and "Project title" columns.')
    code_at, title_at = headers.index(CODE_HEADER), headers.index(TITLE_HEADER)

    found = []
    for number, row in enumerate(rows, start=2):
        code = str(row[code_at] or "").strip() if code_at < len(row) else ""
        title = " ".join(str(row[title_at] or "").split()) if title_at < len(row) else ""
        if code or title:
            found.append((number, code, title))
    return found


class Command(BaseCommand):
    help = "Set the submitted project title of each team listed in a spreadsheet."

    def add_arguments(self, parser):
        parser.add_argument("path", help="An .xlsx file with Team code and Project title columns.")
        parser.add_argument("--year", type=int, required=True, help="The year the teams belong to.")
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Report what would change without saving anything.",
        )

    def handle(self, *args, path, year, dry_run, **options):
        rows = read_rows(path)
        groups = {
            group.group_name.strip().upper(): group
            for group in Groups.objects.active().filter(year=year)
        }
        entries = {
            entry.group_id: entry
            for entry in Submission.objects.filter(group__in=groups.values())
        }

        updated, unchanged, problems, no_entry = [], [], [], []
        seen = set()
        for number, code, title in rows:
            key = code.upper()
            if not code or not title:
                problems.append(f"row {number}: missing a team code or a title")
            elif key in seen:
                problems.append(f"row {number}: {code} is listed more than once")
            elif len(title) > MAX_TITLE_LENGTH:
                problems.append(f"row {number}: {code}'s title is over {MAX_TITLE_LENGTH} characters")
            elif key not in groups:
                problems.append(f"row {number}: no {year} team named {code}")
            elif groups[key].id not in entries:
                no_entry.append(code)
            else:
                entry = entries[groups[key].id]
                if entry.project_title == title and entry.submitted_project_title == title:
                    unchanged.append(code)
                else:
                    entry.project_title = title
                    entry.submitted_project_title = title
                    updated.append(entry)
            seen.add(key)

        if not dry_run:
            with transaction.atomic():
                for entry in updated:
                    # Only the titles, so the entry's status and timestamps stay as they were.
                    entry.save(update_fields=["project_title", "submitted_project_title"])

        for problem in problems:
            self.stdout.write(self.style.ERROR(f"  {problem}"))
        if no_entry:
            self.stdout.write(self.style.WARNING(f"  No entry to attach a title to: {', '.join(no_entry)}"))
        verb = "Would update" if dry_run else "Updated"
        self.stdout.write(
            f"{verb} {len(updated)}. Already up to date: {len(unchanged)}. "
            f"No entry: {len(no_entry)}. Problems: {len(problems)}."
        )
        if dry_run:
            self.stdout.write(self.style.WARNING("Dry run: nothing was saved."))
