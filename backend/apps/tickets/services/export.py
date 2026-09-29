"""The support queue as an Excel file (client item C-09).

Asked for by the client while looking at the queue: "from the admin
perspective, some export function. If we wanted just to export that table."
So the sheet is that table: the same columns under the same headings, filled
from the same row dict the queue endpoint sends (views_admin._queue_row), for
every ticket the current filters match rather than only the page on screen.

ORM-free, like grading/services/xlsx.py: the view decides which tickets are
in the file, and this module only decides how they look in it.
"""

from __future__ import annotations

import io
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from ..models import TicketPriority, TicketStatus

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# The queue table's headings, in its order (adminweb QueueTable.tsx). Its
# first column is the selection checkbox, which holds no data, so it is the
# only one not here. A heading changed on either side should change on both.
HEADERS = [
    "Ticket",
    "Requester",
    "Subject",
    "Status",
    "Priority",
    "Assignee",
    "Last activity",
]

_WIDTHS = [16, 26, 50, 22, 10, 22, 28]

# What the table shows in the Assignee column when nobody owns the ticket.
UNASSIGNED_LABEL = "Unassigned"

# The name line of a ticket with no requester. The table shows a dash there;
# this is a plain hyphen because nothing a person reads may carry an em-dash.
# It starts with one of the formula triggers below, so _put marks the cell as
# text. That is wanted: edited in Excel, "-" over a region would otherwise be
# read as a minus sign in front of it.
NO_REQUESTER_LABEL = "-"

# A cell whose text starts with one of these can be read by a spreadsheet as
# something to run rather than something to show. Subjects and requester names
# are typed by students, so a subject of '=HYPERLINK("http://...","Open")'
# would otherwise reach an admin's Excel as a live link, one click from
# wherever the student chose. The list is the one OWASP gives for CSV
# injection. xlsx is less exposed than CSV, because a text cell is not parsed
# when the file opens; _put says what is still exposed and how it is closed.
_FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r")


def export_zone(name) -> ZoneInfo:
    """The zone the exporter's times are written in: their profile's, else UTC.

    The table shows each reader their own browser's time, and a file has to
    pick one zone for everybody who later opens it. The exporter's profile
    zone is the platform's existing answer to "whose clock" for times the
    server writes out (event reminder emails read the same field). It
    defaults to UTC and most admins have never changed it, which is why every
    time cell states its zone rather than leaving the reader to assume one.

    An unknown or empty value falls back to UTC instead of failing: a bad
    profile row must not stop the export.
    """
    try:
        return ZoneInfo(name or "UTC")
    except (ZoneInfoNotFoundError, ValueError, TypeError):
        return ZoneInfo("UTC")


def _put(ws, row, column, value):
    """Write one cell so that text always reads back as the text it was.

    Two things can stop that, and both are reachable from a student's
    keyboard.

    Control characters: openpyxl refuses \\x00-\\x08, \\x0b, \\x0c and
    \\x0e-\\x1f with IllegalCharacterError, so one ticket carrying a vertical
    tab would be a 500 for every export that includes it. The ticket form
    accepts them (measured: a subject with \\x0b in the middle is a 201 and is
    stored as sent), and the registration view writes names straight from
    the request body. They are dropped, and dropped *first*: stripping after
    the check below would turn "\\x0b=1+1" into a formula the check had
    already waved through.

    Formula triggers: openpyxl stores any string starting with "=" as a
    formula, so the data type is forced back to text. That covers what the
    file holds, not what happens when somebody edits the cell: Excel reads an
    edited cell as if it had just been typed, and "=..." or "+..." typed into
    a cell is a formula. quotePrefix is Excel's own "this is text" marker
    (what typing a leading apostrophe sets), and editing keeps it. It also
    leaves the text exactly as typed, which an apostrophe written into the
    value would not.
    """
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    cell = ws.cell(row=row, column=column, value=value)
    if isinstance(value, str) and value.startswith(_FORMULA_TRIGGERS):
        cell.data_type = "s"
        cell.quotePrefix = True
    return cell


def _requester(row):
    # Name, then region on the line below, as the table lays them out. Tickets
    # the platform raised itself have no requester, and the table still draws
    # the name line, with a dash on it. The line is kept here as well: those
    # tickets carry the flagged sender's region (handoff._region_snapshot),
    # so without it "Australia" would sit where a name goes and read as one.
    name = row["user"]["name"] or NO_REQUESTER_LABEL
    region = row["user"]["region"]
    return f"{name}\n{region}" if region else name


def _status(row):
    # The table puts an Overdue badge beside the status badge, in the same
    # column, so the flag travels in the same cell.
    label = TicketStatus(row["status"]).label
    return f"{label}, Overdue" if row["overdue"] else label


def queue_workbook(rows, *, zone: ZoneInfo) -> bytes:
    """XLSX bytes for ``rows``: the queue endpoint's own row dicts, in order."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Tickets"

    for column, heading in enumerate(HEADERS, start=1):
        ws.cell(row=1, column=column, value=heading).font = Font(bold=True)

    for index, row in enumerate(rows, start=2):
        assignee = row["assignee"]
        values = [
            row["ticketNumber"],
            _requester(row),
            row["subject"],
            _status(row),
            TicketPriority(row["priority"]).label,
            assignee["name"] if assignee else UNASSIGNED_LABEL,
        ]
        for column, value in enumerate(values, start=1):
            _put(ws, index, column, value)

        # A real date, not text, so the column sorts and filters as dates.
        # Excel has no time zones, so the value is the local wall-clock time
        # and the zone is written into the cell's display format, per cell:
        # a Sydney export spans AEST and AEDT, and a single label on the
        # heading would be wrong for half the year.
        local = row["supportUpdatedAt"].astimezone(zone)
        when = ws.cell(row=index, column=len(values) + 1, value=local.replace(tzinfo=None))
        when.number_format = f'd mmm yyyy, hh:mm AM/PM "{local.tzname()}"'

    for column, width in enumerate(_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "A2"
    for sheet_row in ws.iter_rows(min_row=2):
        for cell in sheet_row:
            # Wrapped so the requester's second line shows, and top-aligned so
            # a two-line row does not float its other cells to the middle.
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
