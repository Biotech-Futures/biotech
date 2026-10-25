"""A supervisor's Upload Guardian Details on My Students: a sheet of their
students' guardians, checked row by row, then saved all at once.

Laid out like Grading's Upload marks: the first run previews (``dry_run``),
and Apply parses the file again so it acts on the database as it is then.
Any problem row stops the whole sheet, so nothing is half saved.
"""
from __future__ import annotations

import csv
import io

from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from .serializers import is_placeholder_email, is_student_email

HEADERS = ["student_email", "guardian_first_name", "guardian_last_name", "guardian_email"]


def _read_sheet(file, filename: str) -> tuple[list[str], list[tuple[int, dict]]]:
    """The header and ``(row_number, values)`` rows of a CSV or XLSX, numbered
    as a spreadsheet shows them (the header is row 1). Blank rows are skipped."""
    if (filename or "").lower().endswith(".csv"):
        content = file.read()
        if isinstance(content, bytes):
            content = content.decode("utf-8-sig")
        lines = list(csv.reader(io.StringIO(content)))
    else:
        sheet = load_workbook(file, read_only=True, data_only=True).active
        lines = [] if sheet is None else [
            ["" if value is None else str(value) for value in row]
            for row in sheet.iter_rows(values_only=True)
        ]
    if not lines:
        return [], []
    header = [str(cell).strip() for cell in lines[0]]
    rows = []
    for number, line in enumerate(lines[1:], start=2):
        values = {key: str(value).strip() for key, value in zip(header, line) if key}
        if any(values.values()):
            rows.append((number, values))
    return header, rows


def _is_email(value: str) -> bool:
    try:
        validate_email(value)
    except ValidationError:
        return False
    return True


def parse_guardian_upload(file, filename: str, students) -> dict:
    """What a sheet would change for the students in ``students`` (a
    supervisor's own, or every one for an admin), with every problem found.
    Nothing is saved."""
    result = {
        "checks": {
            "missing_headers": [],
            "header_order": "",
            "unknown_students": [],
            "bad_rows": [],
            "consented": [],
        },
        "creates": [],
        "updates": [],
        "unchanged": [],
    }
    checks = result["checks"]
    header, rows = _read_sheet(file, filename)

    checks["missing_headers"] = [name for name in HEADERS if name not in header]
    if not checks["missing_headers"] and header[: len(HEADERS)] != HEADERS:
        checks["header_order"] = f"Out of order (should be {', '.join(HEADERS)})"
    if checks["missing_headers"] or checks["header_order"]:
        result["summary"] = _summary(result)
        return result

    profiles = {
        (profile.user.email or "").strip().lower(): profile
        for profile in students.select_related("user")
    }
    seen = set()
    for number, values in rows:
        student_email = values.get("student_email", "").lower()
        first = values.get("guardian_first_name", "")
        last = values.get("guardian_last_name", "")
        guardian_email = values.get("guardian_email", "").lower()

        profile = profiles.get(student_email)
        if profile is None:
            checks["unknown_students"].append({"row": number, "email": student_email or "blank"})
            continue
        if student_email in seen:
            checks["bad_rows"].append({"row": number, "reason": "student listed twice"})
            continue
        seen.add(student_email)

        problems = [
            column for column, value in (
                ("guardian_first_name", first), ("guardian_last_name", last), ("guardian_email", guardian_email),
            ) if not value
        ]
        if guardian_email and not _is_email(guardian_email):
            problems.append("guardian_email not an email")
        elif guardian_email and guardian_email == student_email:
            problems.append("guardian_email is the student's")
        # A new guardian only: not a placeholder, nor a student's account.
        elif guardian_email and guardian_email != (profile.pg_email or "").strip().lower():
            if is_placeholder_email(guardian_email):
                problems.append("guardian_email is a placeholder (.invalid)")
            elif is_student_email(guardian_email):
                problems.append("guardian_email is a student's account")
        if problems:
            checks["bad_rows"].append({"row": number, "reason": ", ".join(problems)})
            continue
        if profile.has_join_permission:
            checks["consented"].append({"row": number, "email": student_email})
            continue

        user = profile.user
        entry = {
            "row": number,
            "student_id": profile.user_id,
            "student_name": f"{user.first_name or ''} {user.last_name or ''}".strip() or user.email,
            "guardian_first_name": first,
            "guardian_last_name": last,
            "guardian_email": guardian_email,
        }
        current = (profile.pg_first_name or "", profile.pg_last_name or "", (profile.pg_email or "").strip().lower())
        if current == (first, last, guardian_email):
            result["unchanged"].append(entry)
        elif current[2]:
            # A guardian email on file is a record this replaces.
            result["updates"].append(entry)
        else:
            result["creates"].append(entry)

    result["summary"] = _summary(result)
    return result


def _summary(result: dict) -> dict:
    checks = result["checks"]
    errors = (
        len(checks["missing_headers"]) + bool(checks["header_order"]) + len(checks["unknown_students"])
        + len(checks["bad_rows"]) + len(checks["consented"])
    )
    return {
        "creates": len(result["creates"]),
        "updates": len(result["updates"]),
        "unchanged": len(result["unchanged"]),
        "errors": errors,
    }


def commit_guardian_upload(result: dict, *, students, actor, log_edit) -> int:
    """Save a sheet with no problems for the students in ``students``;
    returns how many changed. ``log_edit`` records each change as ``actor``'s."""
    entries = result["creates"] + result["updates"]
    with transaction.atomic():
        profiles = {
            profile.user_id: profile
            for profile in students.select_for_update().filter(
                user_id__in=[entry["student_id"] for entry in entries],
            )
        }
        now = timezone.now()
        for entry in entries:
            profile = profiles[entry["student_id"]]
            profile.pg_first_name = entry["guardian_first_name"]
            profile.pg_last_name = entry["guardian_last_name"]
            profile.pg_email = entry["guardian_email"]
            profile.parent_guardian_flag = True
            # A supervisor's or admin's edit makes the student's details theirs to change.
            profile.supervisor_edited_at = now
            profile.save(update_fields=[
                "pg_first_name", "pg_last_name", "pg_email", "parent_guardian_flag", "supervisor_edited_at",
            ])
            log_edit(actor, profile.user_id, "guardian_update", {"guardianEmail": entry["guardian_email"]})
    return len(entries)
