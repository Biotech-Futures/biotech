"""Docx generation: both template dialects (tokens and content controls),
plus director names and signature images in rendered documents."""
import io
import zipfile
from decimal import Decimal

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase

from apps.grading.models import ComponentFeedback, Grade, GradingSettings
from apps.grading.services.docx import (
    _fill_content_controls,
    _render_token_template,
    _sum_marks,
)

from .fixtures import _GradingFixture, _build_docx, _seed_doc_templates


class ClientDocxTemplateTests(_GradingFixture):
    """Both template dialects render correctly from uploaded templates.

    The marks release path uses {{Field}} tokens; the certificate path uses
    Word content controls with aliases. Both must come back with placeholders
    replaced and our data in place.
    """

    @staticmethod
    def _document_xml(data: bytes) -> str:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            return z.read("word/document.xml").decode("utf8")

    @staticmethod
    def _control_template(aliases: list[str]) -> bytes:
        """A docx whose only placeholders are content controls with aliases."""
        base = _build_docx("Certificate")
        sdt = "".join(
            f'<w:sdt><w:sdtPr><w:alias w:val="{a}"/></w:sdtPr>'
            f"<w:sdtContent><w:p><w:r><w:t>___</w:t></w:r></w:p></w:sdtContent></w:sdt>"
            for a in aliases
        )
        with zipfile.ZipFile(io.BytesIO(base)) as zin:
            parts = {name: zin.read(name) for name in zin.namelist()}
        xml = parts["word/document.xml"].decode("utf8")
        cut = xml.find("<w:sectPr")
        if cut != -1:
            xml = xml[:cut] + sdt + xml[cut:]
        else:
            xml = xml.replace("</w:body>", sdt + "</w:body>")
        parts["word/document.xml"] = xml.encode("utf8")
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            for name, blob in parts.items():
                zout.writestr(name, blob)
        return out.getvalue()

    def test_marks_release_tokens_filled(self):
        from apps.grading.services.docx import marks_summary_context, render_marks_summary
        from apps.grading.views.student import _grades_payload

        _seed_doc_templates()
        Grade.objects.create(
            submission=self.saq_submission, criterion=self.saq_c1,
            mark=Decimal("4.00"), comment="Nice claim.", graded_by=self.staff,
        )
        Grade.objects.create(
            submission=self.poster_submission, criterion=self.poster_c1,
            mark=Decimal("3.50"), graded_by=self.staff,
        )
        ComponentFeedback.objects.create(
            group=self.group, component=self.poster, comment="Strong poster overall.",
        )
        components = _grades_payload(self.group, 2026)
        data = render_marks_summary(marks_summary_context(self.group, 2026, components))
        xml = self._document_xml(data)
        self.assertNotIn("<<[", xml)
        self.assertNotIn("{{", xml)
        self.assertIn("BTF-TEST-1", xml)      # TeamCode
        self.assertIn("4.00", xml)            # S1 mark
        self.assertIn("Nice claim.", xml)     # S1 comment
        self.assertIn("7.50", xml)            # CombinedTotal = 4.00 + 3.50
        self.assertIn("Strong poster overall.", xml)  # PosterComment

    def test_certificate_content_controls_filled(self):
        from apps.grading.services.docx import (
            certificate_context,
            render_participation_certificate,
        )

        row = GradingSettings.load()
        row.certificate_template = SimpleUploadedFile(
            "cert.docx",
            self._control_template(["firstName", "lastName", "projectTitle"]),
        )
        row.save()

        data = render_participation_certificate(
            certificate_context(
                "Ada Grader", "BTF-TEST-1", 2026, first_name="Ada", last_name="Grader",
            )
        )
        xml = self._document_xml(data)
        self.assertIn("Ada", xml)
        self.assertIn("Grader", xml)
        self.assertIn("BTF-TEST-1", xml)      # projectTitle falls back to group name
        self.assertNotIn("___", xml)          # every control's placeholder replaced


class DocxEngineEdgeTests(SimpleTestCase):
    """The rendering engine's awkward inputs: tokens Word has split across
    runs, tokens inside tables, corrupt images, and unknown control aliases.
    Exercised at the engine seam — no DB needed."""

    @staticmethod
    def _document_xml(data: bytes) -> str:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            return z.read("word/document.xml").decode("utf8")

    def test_a_token_split_across_runs_is_still_replaced(self):
        from docx import Document as NewDocument

        doc = NewDocument()
        p = doc.add_paragraph()
        # Word routinely fragments "{{TeamCode}}" like this after edits.
        p.add_run("{{Team")
        p.add_run("Code}}")
        buf = io.BytesIO()
        doc.save(buf)

        xml = self._document_xml(_render_token_template(buf.getvalue(), {"TeamCode": "BTF-9"}))
        self.assertIn("BTF-9", xml)
        self.assertNotIn("{{", xml)

    def test_tokens_inside_table_cells_are_replaced(self):
        from docx import Document as NewDocument

        doc = NewDocument()
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).paragraphs[0].add_run("Team {{TeamCode}}")
        buf = io.BytesIO()
        doc.save(buf)

        xml = self._document_xml(_render_token_template(buf.getvalue(), {"TeamCode": "BTF-9"}))
        self.assertIn("BTF-9", xml)
        self.assertNotIn("{{TeamCode}}", xml)

    def test_a_corrupt_signature_image_does_not_sink_the_document(self):
        data = _build_docx("{{Director1Signature}} {{Director1Name}}")
        rendered = _render_token_template(
            data,
            {"Director1Name": "Prof. Alice Adams"},
            images={"Director1Signature": b"plainly not pixels"},
        )
        with zipfile.ZipFile(io.BytesIO(rendered)) as z:
            names = z.namelist()
            xml = z.read("word/document.xml").decode("utf8")
        # No picture landed, the token is cleared, and the name still prints.
        self.assertFalse([n for n in names if n.startswith("word/media/")], names)
        self.assertNotIn("{{", xml)
        self.assertIn("Prof. Alice Adams", xml)

    def test_an_unknown_content_control_alias_is_left_alone(self):
        data = ClientDocxTemplateTests._control_template(["firstName", "NotAField"])
        xml = self._document_xml(_fill_content_controls(data, {"firstName": "Ada"}))
        self.assertIn("Ada", xml)
        # The unrecognised control keeps its placeholder rather than being
        # silently blanked — visible in the output means fixable by an admin.
        self.assertIn("___", xml)

    def test_sum_marks_skips_unparseable_values(self):
        total = _sum_marks([
            {"mark": "3.50"},
            {"mark": "abc"},
            {"mark": None},
            {"mark": "1.25"},
        ])
        self.assertEqual(total, Decimal("4.75"))


class DirectorSignatureTests(_GradingFixture):
    """Director names and signature images reach the rendered documents.

    The client's own templates carry no director placeholders, so these build
    a token template on the fly — the same path an admin gets after adding the
    tokens to their docx.
    """

    # Smallest valid PNG: python-docx reads its real dimensions from IHDR.
    PNG_1PX = (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01"
        b"\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
    )

    @staticmethod
    def _token_template(text: str) -> bytes:
        from docx import Document as NewDocument

        doc = NewDocument()
        doc.add_paragraph(text)
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    def _configure(self, template_text: str, *, with_signatures: bool):
        from apps.grading.models import GradingSettings

        settings_row = GradingSettings.load()
        settings_row.director_1_name = "Prof. Alice Adams"
        settings_row.director_2_name = "Dr. Bob Brown"
        settings_row.marks_summary_template = SimpleUploadedFile(
            "tpl.docx", self._token_template(template_text)
        )
        if with_signatures:
            settings_row.director_1_signature = SimpleUploadedFile("sig1.png", self.PNG_1PX)
            settings_row.director_2_signature = SimpleUploadedFile("sig2.png", self.PNG_1PX)
        settings_row.save()
        return settings_row

    def _render(self):
        from apps.grading.services.docx import marks_summary_context, render_marks_summary

        return render_marks_summary(marks_summary_context(self.group, 2026, []))

    def test_director_names_fill_their_tokens(self):
        self._configure(
            "Signed: {{Director1Name}} and {{Director2Name}}", with_signatures=False
        )
        with zipfile.ZipFile(io.BytesIO(self._render())) as z:
            xml = z.read("word/document.xml").decode("utf8")
        self.assertIn("Prof. Alice Adams", xml)
        self.assertIn("Dr. Bob Brown", xml)
        self.assertNotIn("{{", xml)

    def test_signature_tokens_become_images(self):
        self._configure(
            "{{Director1Signature}} {{Director2Signature}} {{Director1Name}}",
            with_signatures=True,
        )
        payload = self._render()
        with zipfile.ZipFile(io.BytesIO(payload)) as z:
            names = z.namelist()
            xml = z.read("word/document.xml").decode("utf8")
        # Two inline pictures, and the token text gone. Only one media part is
        # expected: both fixtures are byte-identical, so python-docx stores the
        # image once and both drawings point at it.
        self.assertEqual(xml.count("<w:drawing>"), 2, xml[:400])
        self.assertTrue([n for n in names if n.startswith("word/media/")], names)
        self.assertNotIn("{{", xml)
        self.assertIn("Prof. Alice Adams", xml)

    def test_missing_signature_leaves_document_renderable(self):
        # No signature uploaded: the token clears and the name still prints.
        self._configure("{{Director1Signature}}{{Director1Name}}", with_signatures=False)
        with zipfile.ZipFile(io.BytesIO(self._render())) as z:
            names = z.namelist()
            xml = z.read("word/document.xml").decode("utf8")
        self.assertFalse([n for n in names if n.startswith("word/media/")], names)
        self.assertNotIn("{{", xml)
        self.assertIn("Prof. Alice Adams", xml)
