"""Export downloads: the sync per-group zip, the async per-component job,
the job polling endpoint, and the SAQ XLSX export round-trip."""
import io
import zipfile
from decimal import Decimal
from unittest import mock

from django.test import SimpleTestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from rest_framework import status
from rest_framework.test import APIClient

from apps.grading.models import Grade, GradingJob, GroupMarkingCategories
from apps.grading.services import zip as zip_service
from apps.grading.services.content import ComponentEntry
from apps.grading.services.xlsx import build_saq_xlsx

from .fixtures import _GradingFixture


class BuildSubmissionsZipTests(SimpleTestCase):
    """Direct coverage of the zip builder's concurrent blob prefetch: with
    far more entries than pool workers, the archive must still come out in
    entry order with each entry's own bytes — byte-identical to what the old
    sequential loop produced."""

    @staticmethod
    def _entry(i: int, *, file: dict | None = None, text: str = "", link: str = "") -> ComponentEntry:
        return ComponentEntry(
            submission_id=i, group_id=i, group_name=f"Group-{i:03d}",
            component_id=2, component_code="POSTER",
            submitted_at=None, is_late=False,
            file=file, text=text, link=link,
        )

    def test_prefetch_preserves_order_content_and_missing_markers(self):
        blobs = {f"key-{i}": f"payload-{i}".encode() for i in range(40)}
        entries = []
        for i in range(40):
            entries.append(self._entry(i, file={"storage_key": f"key-{i}", "name": f"p{i}.pdf"}))
            # Interleave file-less entries: they hold no pool slot and must
            # not disturb the ordering around them.
            if i % 10 == 0:
                entries.append(self._entry(1000 + i, text="answers", link="https://x.example"))
        entries.append(self._entry(999, file={"storage_key": "gone", "name": "lost.pdf"}))

        def fake_open(entry):
            key = entry.file["storage_key"]
            if key not in blobs:
                raise FileNotFoundError(key)
            return io.BytesIO(blobs[key])

        with mock.patch.object(zip_service, "open_file", side_effect=fake_open):
            payload = zip_service.build_submissions_zip(entries, group_folder=False)

        zf = zipfile.ZipFile(io.BytesIO(payload))
        year = timezone.now().year
        expected = []
        for entry in entries:
            stem = f"{year}_{entry.group_name}_Poster"
            if entry.file and entry.file["storage_key"] in blobs:
                expected.append(f"{stem}.pdf")
            elif entry.file:
                expected.append(f"{stem}_MISSING.txt")
            else:
                expected.extend([f"{stem}.txt", f"{stem}_Link.txt"])
        self.assertEqual(zf.namelist(), expected)
        for i in range(40):
            self.assertEqual(zf.read(f"{year}_Group-{i:03d}_Poster.pdf"), blobs[f"key-{i}"])
        self.assertIn(b"Original blob missing: gone", zf.read(f"{year}_Group-999_Poster_MISSING.txt"))


class GroupDownloadViewTests(_GradingFixture):
    """Sync per-group zip. Bounded selection, streamed straight to the client."""

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.staff)

    def test_zip_contains_submission_text_and_link(self):
        # Add a prototype link to the entry and resubmit so the snapshot has it.
        self.submission.prototype_url = "https://example.com/prototype"
        self.submission.snapshot(self.staff)
        self.submission.save()

        url = reverse("grading:group-download", kwargs={"group_id": self.group.id})
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp["Content-Type"], "application/zip")
        self.assertIn("attachment;", resp["Content-Disposition"])

        zf = zipfile.ZipFile(io.BytesIO(resp.content))
        names = set(zf.namelist())
        # Single-group download: no group subfolder, flat
        # <Year>_<Group>_<Component> files at the archive root.
        year = timezone.now().year
        base = f"{year}_BTF-TEST-1"
        self.assertIn(f"{base}_SAQs.txt", names)
        self.assertIn(f"{base}_Prototype_Link.txt", names)
        # The fixture's poster storage key has no backing blob; the archive
        # notes it instead of failing.
        self.assertIn(f"{base}_Poster_MISSING.txt", names)
        self.assertIn("Some student answers.", zf.read(f"{base}_SAQs.txt").decode())
        self.assertEqual(
            zf.read(f"{base}_Prototype_Link.txt").decode().strip(),
            "https://example.com/prototype",
        )

    def test_component_filter(self):
        url = reverse("grading:group-download", kwargs={"group_id": self.group.id}) + "?component=SAQ"
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        names = zipfile.ZipFile(io.BytesIO(resp.content)).namelist()
        year = timezone.now().year
        self.assertEqual(names, [f"{year}_BTF-TEST-1_SAQs.txt"])

    def test_non_staff_denied(self):
        self.client.force_authenticate(self.non_staff)
        resp = self.client.get(reverse("grading:group-download", kwargs={"group_id": self.group.id}))
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)


@override_settings(GRADING_JOB_DISPATCH_SYNC=True)
class ComponentDownloadViewTests(_GradingFixture):
    """Async per-component export. DISPATCH_SYNC runs inline so the job row is
    already ``done`` (or ``failed``) by the time the 202 returns."""

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.staff)

    def test_zip_job_runs_inline_and_produces_url(self):
        url = reverse("grading:component-download", kwargs={"code": "SAQ"})
        resp = self.client.post(url, {"format": "zip"}, format="json")
        self.assertEqual(resp.status_code, status.HTTP_202_ACCEPTED, resp.content)
        job_id = resp.json()["job_id"]
        job = GradingJob.objects.get(pk=job_id)
        self.assertEqual(job.status, GradingJob.STATUS_DONE, job.error)
        self.assertTrue(job.result_url)

    def test_xlsx_only_for_saq(self):
        url = reverse("grading:component-download", kwargs={"code": "POSTER"})
        resp = self.client.post(url, {"format": "xlsx"}, format="json")
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)

    def test_bad_format_rejected(self):
        url = reverse("grading:component-download", kwargs={"code": "SAQ"})
        resp = self.client.post(url, {"format": "docx"}, format="json")
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)

    def test_non_staff_denied(self):
        self.client.force_authenticate(self.non_staff)
        url = reverse("grading:component-download", kwargs={"code": "SAQ"})
        resp = self.client.post(url, {"format": "zip"}, format="json")
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)


@override_settings(GRADING_JOB_DISPATCH_SYNC=True)
class GradingJobDetailViewTests(_GradingFixture):
    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.staff)

    def test_poll_shape(self):
        job = GradingJob.objects.create(
            kind=GradingJob.KIND_BULK_ZIP,
            status=GradingJob.STATUS_DONE,
            params={"kind": "component_zip", "component_code": "SAQ"},
            result_url="grading/jobs/1/SAQ-bundle.zip",  # opaque storage key now
            created_by=self.staff,
        )
        resp = self.client.get(reverse("grading:job-detail", kwargs={"pk": job.pk}))
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        data = resp.json()
        self.assertEqual(data["status"], "done")
        self.assertTrue(data["download_url"].endswith(f"/api/v1/grading/jobs/{job.pk}/download/"))

    def test_unknown_job_404(self):
        resp = self.client.get(reverse("grading:job-detail", kwargs={"pk": 999_999}))
        self.assertEqual(resp.status_code, status.HTTP_404_NOT_FOUND)


@override_settings(GRADING_JOB_DISPATCH_SYNC=True)
class SaqXlsxExportTests(_GradingFixture):
    """The XLSX export is the client's off-platform marking template: the
    sheet must carry existing marks pre-filled, and the bulk-upload parser
    must accept the exact bytes it produces back without a single diff."""

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.staff)
        Grade.objects.create(
            submission=self.saq_submission, criterion=self.saq_c1,
            mark=Decimal("8.00"), comment="Great claim.",
        )

    def _export_xlsx(self) -> bytes:
        resp = self.client.post(
            reverse("grading:component-download", kwargs={"code": "SAQ"}),
            {"format": "xlsx"}, format="json",
        )
        self.assertEqual(resp.status_code, status.HTTP_202_ACCEPTED, resp.content)
        job = GradingJob.objects.get(pk=resp.json()["job_id"])
        self.assertEqual(job.status, GradingJob.STATUS_DONE, job.error)
        # Fetch through the download endpoint, as a marker's browser would.
        download = self.client.get(
            reverse("grading:job-download", kwargs={"pk": job.pk})
        )
        self.assertEqual(download.status_code, status.HTTP_200_OK)
        if download.streaming:
            return b"".join(download.streaming_content)
        return download.content

    def test_sheet_carries_the_upload_shape_with_marks_prefilled(self):
        GroupMarkingCategories.objects.create(
            group=self.group,
            product_categories=["Health", "Other"],
            product_category_other="Wearables",
            solution_category="Other",
            solution_category_other="App",
        )
        ws = load_workbook(io.BytesIO(self._export_xlsx())).active
        rows = list(ws.iter_rows(values_only=True))
        # One answered question and a two-criterion rubric.
        self.assertEqual(
            list(rows[0]),
            ["group_id", "group_name", "type", "q1",
             "r1_mark", "r1_comment", "r2_mark", "r2_comment",
             "overall_comment", "product_category", "category_of_solution"],
        )
        # One row per group.
        self.assertEqual(len(rows), 2)
        (group_id, group_name, row_type, q1, r1_mark, r1_comment, r2_mark, r2_comment,
         overall_comment, product_category, category_of_solution) = rows[1]
        self.assertEqual(group_id, self.group.id)
        self.assertEqual(group_name, "BTF-TEST-1")
        self.assertEqual(row_type, "SAQs")
        # The answer cell carries the answer under its question prompt.
        self.assertIn("Team answers", q1)
        self.assertIn("Some student answers.", q1)
        # Existing grade pre-filled; the ungraded criterion stays blank.
        self.assertEqual(r1_mark, 8.0)
        self.assertEqual(r1_comment, "Great claim.")
        self.assertIsNone(r2_mark)
        self.assertIn(r2_comment, (None, ""))
        # No SAQ feedback saved in this fixture -> blank, not an error.
        self.assertIn(overall_comment, (None, ""))
        # The marking key's selections, with Other's text written plainly in
        # its place (no "Other:" prefix).
        self.assertEqual(product_category, "Health, Wearables")
        self.assertEqual(category_of_solution, "App")

    def test_export_round_trips_through_bulk_upload_without_a_diff(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        upload = SimpleUploadedFile(
            "saq-export.xlsx", self._export_xlsx(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        resp = self.client.post(
            reverse("grading:component-bulk-upload", kwargs={"code": "SAQ"}),
            {"file": upload, "dry_run": "true"},
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK, resp.content)
        body = resp.json()
        self.assertEqual(body["checks"]["missing_headers"], [])
        # The untouched export must read as exactly what is already stored:
        # the pre-filled mark is unchanged, the blank criterion is untouched,
        # and nothing is created, updated or refused.
        self.assertEqual(
            body["summary"],
            {"creates": 0, "updates": 0, "unchanged": 1,
             "overall_comments": 0, "marking_categories": 0, "errors": 0},
        )

    def test_categories_round_trip_without_the_other_prefix(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        # Other text with a comma in it, and a bare Other with no text.
        GroupMarkingCategories.objects.create(
            group=self.group,
            product_categories=["Health and Medicine", "Other"],
            product_category_other="Wearables, apps",
            solution_category="Other",
            solution_category_other="",
        )
        payload = self._export_xlsx()
        rows = list(load_workbook(io.BytesIO(payload)).active.iter_rows(values_only=True))
        self.assertEqual(rows[1][-2:], ("Health and Medicine, Wearables, apps", "Other"))

        upload = SimpleUploadedFile(
            "saq-export.xlsx", payload,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        resp = self.client.post(
            reverse("grading:component-bulk-upload", kwargs={"code": "SAQ"}),
            {"file": upload, "dry_run": "true"},
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK, resp.content)
        self.assertEqual(resp.json()["summary"]["marking_categories"], 0)


class SaqXlsxQuestionColumnsTests(SimpleTestCase):
    """The qN columns line up by question across groups, in form order: a
    group that skipped a question gets a blank cell, never a shifted row."""

    @staticmethod
    def _entry(i: int, answers: list[tuple[str, str]]) -> ComponentEntry:
        return ComponentEntry(
            submission_id=i, group_id=i, group_name=f"BTF{i:02d}",
            component_id=1, component_code="SAQ",
            submitted_at=None, is_late=False,
            file=None, text="", link="", answers=tuple(answers),
        )

    def test_a_skipped_question_leaves_a_blank_cell(self):
        entries = [
            self._entry(1, [("Q one", "a1"), ("Q two", "a2")]),
            self._entry(2, [("Q two", "b2")]),  # skipped Q one
            self._entry(3, [("Q one", "c1"), ("retired_key", "c9")]),
        ]
        payload = build_saq_xlsx(
            entries, questions=["Q one", "Q two", "Q nobody answered"]
        )
        rows = list(load_workbook(io.BytesIO(payload)).active.iter_rows(values_only=True))
        # Known questions in form order, then the retired one; a question
        # nobody answered gets no column.
        self.assertEqual(
            list(rows[0]),
            ["group_id", "group_name", "type", "q1", "q2", "q3",
             "overall_comment", "product_category", "category_of_solution"],
        )
        answers = [[cell or None for cell in row[3:6]] for row in rows[1:]]
        self.assertEqual(answers, [
            ["Q one\na1", "Q two\na2", None],
            [None, "Q two\nb2", None],
            ["Q one\nc1", None, "retired_key\nc9"],
        ])

    def test_layout_bold_questions_top_aligned_rows_fit_content(self):
        payload = build_saq_xlsx(
            [self._entry(1, [("Q one", "a1")]), self._entry(2, [("Q one", "b1")])],
            questions=["Q one"],
        )
        ws = load_workbook(io.BytesIO(payload), rich_text=True).active
        # The question is bold; the answer below it is plain.
        question, answer = ws["D2"].value
        self.assertEqual(question.text, "Q one")
        self.assertTrue(question.font.b)
        self.assertEqual(answer, "\na1")
        # Every cell, the header included, sits at the top of its row.
        self.assertEqual(
            {cell.alignment.vertical for row in ws.iter_rows() for cell in row}, {"top"}
        )
        # Question columns are 43 wide.
        self.assertEqual(ws.column_dimensions["D"].width, 43)
        # Comment columns (rN_comment, overall_comment) and both category
        # columns are 30 wide; the group name keeps the default width.
        headers = [cell.value for cell in ws[1]]
        for index, header in enumerate(headers, start=1):
            letter = get_column_letter(index)
            if header.endswith("comment") or header in ("product_category", "category_of_solution"):
                self.assertEqual(ws.column_dimensions[letter].width, 30)
        self.assertNotIn("B", ws.column_dimensions)
        # No fixed row heights, so Excel fits each row to its tallest cell.
        self.assertIsNone(ws.row_dimensions[2].height)
        self.assertIsNone(ws.row_dimensions[3].height)
