"""Results emails from the Release Results tab: who gets which email, the
client's wording, the files they carry, the survey details, preview, and
sending each audience in batches."""
import io
import zipfile
from datetime import timedelta
from decimal import Decimal
from email.mime.base import MIMEBase
from unittest import mock

from django.core import mail
from django.test import override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from apps.grading.models import (
    CertificatesRelease,
    FinalistFlag,
    Grade,
    GradingSettings,
    MarksRelease,
    ResultsEmailSettings,
    ResultsSupervisorEmail,
    ResultsTeamEmail,
    Rubric,
    RubricCriterion,
    SubmissionComponent,
)
from apps.grading.services.finalist_notify import symposium_today
from apps.grading.services.results_notify import send_results_batch
from apps.groups.models import GroupMembership, Groups
from apps.services.models import SystemEmailTemplate
from apps.submissions.models import Submission
from apps.users.models import StudentProfile, SupervisorProfile, User

from .fixtures import _GradingFixture, _seed_doc_templates

LOCMEM = "django.core.mail.backends.locmem.EmailBackend"
SEND = "grading:results-email-send"


def _user(email, first="", last=""):
    return User.objects.create_user(email=email, first_name=first, last_name=last, password="pw12345!")


def _student(email, group, supervisor_profile=None):
    user = _user(email, "Stu", email.split("@")[0])
    StudentProfile.objects.create(
        user=user, pg_first_name="P", pg_last_name="G", parent_guardian_flag=True,
        school_name="Test School", year_lvl="10", supervisor=supervisor_profile,
    )
    GroupMembership.objects.create(user=user, group=group, membership_role="student")
    return user


def _files(message):
    """A sent email's files by name (the inline logo left out)."""
    return {name: content for name, content, _ in (a for a in message.attachments if not isinstance(a, MIMEBase))}


def _docx_text(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return z.read("word/document.xml").decode("utf8")


def _submitted_team(name, staff):
    team = Groups.objects.create(group_name=name)
    submission = Submission.objects.create(group=team, answers={"q_answers": "Answers."})
    submission.snapshot(staff)
    submission.save()
    return team


@override_settings(EMAIL_BACKEND=LOCMEM)
class ResultsEmailTests(_GradingFixture):
    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.staff)

        # The fixture's submitted group: two students, a mentor, and a supervisor
        # who is also a member. The students and the mentor get the group email.
        self.supervisor = _user("sam.lee@example.com", "Sam", "Lee")
        self.sup_profile = SupervisorProfile.objects.create(user=self.supervisor)
        _student("amy@example.com", self.group, self.sup_profile)
        _student("ben@example.com", self.group, self.sup_profile)
        mentor = _user("mentor@example.com", "Mo", "Mentor")
        GroupMembership.objects.create(user=mentor, group=self.group, membership_role="mentor")
        GroupMembership.objects.create(user=self.supervisor, group=self.group, membership_role="supervisor")

        # A finalist team (left out while certificates go out without finalists)
        # and a team that never submitted.
        self.finalist = _submitted_team("Finalist Team", self.staff)
        _student("fin@example.com", self.finalist, self.sup_profile)
        FinalistFlag.objects.create(group=self.finalist, flagged_by=self.staff)
        no_entry = Groups.objects.create(group_name="No Entry")
        _student("none@example.com", no_entry)

        now = timezone.now()
        marks = MarksRelease.load()
        marks.released_at = now
        marks.save()
        certificates = CertificatesRelease.load()
        certificates.released_at = now
        certificates.exclude_finalists = True
        certificates.save()
        details = ResultsEmailSettings.load()
        details.survey_closes = symposium_today() + timedelta(days=30)
        details.save()
        _seed_doc_templates()
        Grade.objects.create(
            submission=self.submission, criterion=self.saq_c1, mark=Decimal("8.00"), graded_by=self.staff,
        )
        Grade.objects.create(
            submission=self.submission, criterion=self.poster_c1, mark=Decimal("6.50"), graded_by=self.staff,
        )

    def _send_all(self, audience):
        cursor, results = None, []
        while True:
            r = self.client.post(reverse(SEND), {"audience": audience, "cursor": cursor}, format="json")
            self.assertEqual(r.status_code, status.HTTP_200_OK, r.content)
            results.append(r.json())
            cursor = r.json()["cursor"]
            if r.json()["done"]:
                return results

    def _recipients(self):
        return sorted(m.to[0] for m in mail.outbox)

    # -- who gets it ------------------------------------------------------------

    def test_the_groups_students_and_mentors_get_the_group_email(self):
        results = self._send_all("groups")
        # The supervisor who is a member gets the supervisor email instead.
        self.assertEqual(self._recipients(), ["amy@example.com", "ben@example.com", "mentor@example.com"])

        message = mail.outbox[0]
        text = " ".join(message.body.split())  # the plain text wraps its lines
        self.assertEqual(message.subject, f"Your {self.group.year} BIOTech Futures Challenge results")
        self.assertIn(f"Dear {self.group.group_name},", text)
        self.assertIn(f"Congratulations on your participation in the {self.group.year} BIOTech Futures Challenge!", text)
        self.assertIn("https://sydney.au1.qualtrics.com/jfe/form/SV_cCKb80Gg7IhgBpA", text)
        closes = ResultsEmailSettings.load().survey_closes
        self.assertIn(f"until the {closes.day}", text)
        self.assertIn(f" of {closes:%B}.", text)
        self.assertEqual(message.reply_to, ["support@biotechfutures.org"])

        self.assertEqual(results[-1]["emailed"], 3)
        self.assertEqual(results[-1]["groups"], {"total": 1, "emailed": 1})
        self.assertEqual(results[-1]["supervisors"], {"total": 1, "emailed": 0})
        self.assertTrue(ResultsTeamEmail.objects.filter(group=self.group).exists())

    def test_supervisors_are_emailed_separately_once_each(self):
        results = self._send_all("supervisors")
        self.assertEqual(self._recipients(), ["sam.lee@example.com"])
        message = mail.outbox[0]
        self.assertEqual(message.subject, f"Your students’ {self.group.year} BIOTech Futures Challenge results")
        self.assertIn("Dear Sam Lee,", message.body)
        self.assertIn("Merit certificates for each of your students", " ".join(message.body.split()))
        self.assertEqual(message.reply_to, ["support@biotechfutures.org"])
        self.assertEqual(results[-1]["supervisors"], {"total": 1, "emailed": 1})
        self.assertEqual(results[-1]["groups"], {"total": 1, "emailed": 0})
        self.assertTrue(ResultsSupervisorEmail.objects.filter(supervisor=self.supervisor).exists())

    def test_finalists_are_included_when_certificates_go_to_everyone(self):
        CertificatesRelease.objects.update(exclude_finalists=False)
        self._send_all("groups")
        self._send_all("supervisors")
        self.assertIn("fin@example.com", self._recipients())
        # Their supervisor is the same person: still one supervisor email.
        self.assertEqual(self._recipients().count("sam.lee@example.com"), 1)

    def test_a_second_send_emails_nobody_again(self):
        self._send_all("groups")
        self._send_all("supervisors")
        mail.outbox = []
        self.assertTrue(self._send_all("groups")[-1]["done"])
        self.assertTrue(self._send_all("supervisors")[-1]["done"])
        self.assertEqual(mail.outbox, [])

    # -- the files they carry ------------------------------------------------------

    def test_everyone_in_the_group_gets_all_its_certificates_and_the_marks_summary(self):
        self._send_all("groups")
        year = self.group.year
        self.assertEqual(len(mail.outbox), 3)
        for message in mail.outbox:
            # Students and mentor alike: each sees the others' certificates.
            files = _files(message)
            self.assertEqual(list(files), [
                f"{year}_BTF_Certificate_Stu_amy.docx",
                f"{year}_BTF_Certificate_Stu_ben.docx",
                f"{year}_BTF_Mentor_Certificate_Mo_Mentor.docx",
                f"{year}_BTF_Marks_BTF-TEST-1.docx",
            ])
            self.assertIn("Stu amy — BTF-TEST-1", _docx_text(files[f"{year}_BTF_Certificate_Stu_amy.docx"]))
            self.assertIn(
                "Mo Mentor mentored BTF-TEST-1", _docx_text(files[f"{year}_BTF_Mentor_Certificate_Mo_Mentor.docx"]),
            )
            self.assertIn("Team BTF-TEST-1 S1 8 ", _docx_text(files[f"{year}_BTF_Marks_BTF-TEST-1.docx"]))

    def test_the_files_come_after_the_email_and_its_logo(self):
        self._send_all("groups")
        top = mail.outbox[0].message()
        self.assertEqual(top.get_content_type(), "multipart/mixed")
        body, *files = top.get_payload()
        self.assertEqual(body.get_content_type(), "multipart/related")
        self.assertEqual(
            [part.get_content_type() for part in body.get_payload()],
            ["multipart/alternative", "image/png"],
        )
        self.assertEqual([f.get_content_disposition() for f in files], ["attachment"] * 4)

    def test_supervisors_get_their_students_certificates_and_a_marks_sheet(self):
        # Only Poster and SAQ, as the email says: another marked component stays out.
        report = SubmissionComponent.objects.exclude(code__in=("SAQ", "POSTER")).first()
        RubricCriterion.objects.create(
            rubric=Rubric.objects.create(component=report, year=2026, active=True),
            name="Structure", max_mark=Decimal("5.00"), order=10,
        )
        self._send_all("supervisors")
        year = self.group.year
        files = _files(mail.outbox[0])
        # Their finalist student is left out while certificates exclude finalists.
        self.assertEqual(list(files), [
            f"{year}_BTF_Certificate_Stu_amy.docx",
            f"{year}_BTF_Certificate_Stu_ben.docx",
            f"{year}_BTF_Student_Marks.xlsx",
        ])
        sheet = load_workbook(io.BytesIO(files[f"{year}_BTF_Student_Marks.xlsx"])).active
        rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
        self.assertEqual(rows[0], [
            "Student", "Team",
            "Short Answer Questions: Content (/10)", "Short Answer Questions: Clarity (/5)",
            "Short Answer Questions Total (/15)",
            "A2 Poster: Design (/10)", "A2 Poster Total (/10)",
            "Total (/25)",
        ])
        self.assertEqual(rows[1:], [
            ["Stu amy", "BTF-TEST-1", 8, None, 8, 6.5, 6.5, 14.5],
            ["Stu ben", "BTF-TEST-1", 8, None, 8, 6.5, 6.5, 14.5],
        ])

    def test_a_team_whose_files_cant_be_made_stays_pending(self):
        with mock.patch(
            "apps.grading.services.results_notify.Documents.marks_summary", side_effect=OSError("storage down"),
        ), self.assertLogs("apps.grading.services.results_notify", level="ERROR"):
            results = self._send_all("groups")
        self.assertEqual(mail.outbox, [])
        self.assertEqual(results[-1]["failed"], 1)
        self.assertFalse(ResultsTeamEmail.objects.filter(group=self.group).exists())

    def test_sending_waits_for_the_templates_its_files_need(self):
        GradingSettings.objects.update(mentor_certificate_template="")
        r = self.client.post(reverse(SEND), {"audience": "groups"}, format="json")
        self.assertEqual(
            r.json()["detail"],
            "Upload the marks summary, student certificate and mentor certificate templates "
            "in Document Setup before emailing groups.",
        )
        # The supervisors' spreadsheet needs no template.
        self._send_all("supervisors")
        self.assertEqual(self._recipients(), ["sam.lee@example.com"])

        GradingSettings.objects.update(certificate_template="")
        r = self.client.post(reverse(SEND), {"audience": "supervisors"}, format="json")
        self.assertEqual(
            r.json()["detail"], "Upload the student certificate template in Document Setup before emailing supervisors.",
        )
        r = self.client.get(reverse("grading:results-email"))
        self.assertEqual(r.json()["templates_ready"], {"groups": False, "supervisors": False})

    # -- batches and failures ------------------------------------------------------

    def test_batches_move_on_and_report_progress(self):
        for n in range(6):
            team = _submitted_team(f"Extra {n}", self.staff)
            _student(f"extra{n}@example.com", team)
        first = send_results_batch(self.staff, "groups", None, limit=5)
        # The fixture group has two students and a mentor.
        self.assertEqual((first["emailed"], first["done"]), (7, False))
        second = send_results_batch(self.staff, "groups", first["cursor"], limit=5)
        self.assertEqual((second["emailed"], second["done"]), (2, True))
        self.assertEqual(second["groups"], {"total": 7, "emailed": 7})

    def test_a_team_whose_student_misses_it_stays_pending_and_is_not_retried_in_the_run(self):
        real_send = mail.EmailMultiAlternatives.send

        def fail_for_ben(message, *args, **kwargs):
            if message.to == ["ben@example.com"]:
                raise OSError("rejected")
            return real_send(message, *args, **kwargs)

        with mock.patch("django.core.mail.EmailMultiAlternatives.send", autospec=True, side_effect=fail_for_ben), \
                self.assertLogs("apps.grading.services.results_notify", level="ERROR"):
            results = self._send_all("groups")
        self.assertEqual(sum(r["failed"] for r in results), 1)
        self.assertFalse(ResultsTeamEmail.objects.filter(group=self.group).exists())
        self.assertEqual(results[-1]["groups"], {"total": 1, "emailed": 0})

        # The next press reaches the team again.
        mail.outbox = []
        self._send_all("groups")
        self.assertTrue(ResultsTeamEmail.objects.filter(group=self.group).exists())

    # -- when it may be sent ----------------------------------------------------------

    def test_sending_waits_for_both_releases(self):
        CertificatesRelease.objects.update(released_at=None)
        for audience in ("groups", "supervisors"):
            r = self.client.post(reverse(SEND), {"audience": audience}, format="json")
            self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)
            self.assertEqual(r.json()["detail"], "Release both marks and certificates before sending the results emails.")
        self.assertEqual(mail.outbox, [])

    def test_only_the_group_email_waits_for_the_survey_details(self):
        ResultsEmailSettings.objects.update(survey_closes=None)
        r = self.client.post(reverse(SEND), {"audience": "groups"}, format="json")
        self.assertEqual(r.json()["detail"], "Set the feedback survey link and close date before emailing groups.")
        ResultsEmailSettings.objects.update(survey_closes=symposium_today() - timedelta(days=1))
        r = self.client.post(reverse(SEND), {"audience": "groups"}, format="json")
        self.assertEqual(r.json()["detail"], "The survey close date can't be before today. Update it before emailing groups.")
        # Supervisors aren't told about the survey, so they can go.
        self._send_all("supervisors")
        self.assertEqual(self._recipients(), ["sam.lee@example.com"])

    def test_a_switched_off_email_is_refused(self):
        SystemEmailTemplate.objects.create(key="results_supervisor", is_enabled=False)
        r = self.client.post(reverse(SEND), {"audience": "supervisors"}, format="json")
        self.assertEqual(r.json()["detail"], "Results: supervisors is switched off on System Emails.")
        r = self.client.get(reverse("grading:results-email"))
        self.assertEqual(r.json()["emails_on"], {"groups": True, "supervisors": False})

    def test_an_unknown_audience_is_refused(self):
        r = self.client.post(reverse(SEND), {"audience": "mentors"}, format="json")
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)

    def test_graders_only(self):
        self.client.force_authenticate(self.non_staff)
        for name in ("grading:results-email", SEND):
            self.assertEqual(self.client.post(reverse(name), {}, format="json").status_code, status.HTTP_403_FORBIDDEN)

    # -- details and preview -------------------------------------------------------------

    def test_details_are_saved_and_a_past_close_date_refused(self):
        closes = symposium_today() + timedelta(days=10)
        r = self.client.patch(
            reverse("grading:results-email"),
            {"survey_url": "https://example.com/survey", "survey_closes": closes.isoformat()},
            format="json",
        )
        self.assertEqual(r.status_code, status.HTTP_200_OK, r.content)
        body = r.json()
        self.assertEqual((body["survey_url"], body["survey_closes"], body["complete"]), ("https://example.com/survey", closes.isoformat(), True))
        self.assertEqual((body["marks_released"], body["certificates_released"]), (True, True))
        self.assertEqual(body["groups"], {"total": 1, "emailed": 0})
        self.assertEqual(body["supervisors"], {"total": 1, "emailed": 0})

        r = self.client.patch(
            reverse("grading:results-email"),
            {"survey_closes": (symposium_today() - timedelta(days=1)).isoformat()},
            format="json",
        )
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)

    def test_preview_shows_each_email_with_unsaved_details(self):
        r = self.client.post(
            reverse("grading:results-email-preview"),
            {"audience": "groups", "survey_url": "https://example.com/draft"},
            format="json",
        )
        self.assertEqual(r.status_code, status.HTTP_200_OK, r.content)
        self.assertEqual(r.json()["to"], self.group.group_name)
        self.assertIn('href="https://example.com/draft"', r.json()["html"])
        self.assertEqual(ResultsEmailSettings.load().survey_url, "https://sydney.au1.qualtrics.com/jfe/form/SV_cCKb80Gg7IhgBpA")

        year = self.group.year
        self.assertEqual(r.json()["attachments"], [
            f"{year}_BTF_Certificate_Stu_amy.docx",
            f"{year}_BTF_Certificate_Stu_ben.docx",
            f"{year}_BTF_Mentor_Certificate_Mo_Mentor.docx",
            f"{year}_BTF_Marks_BTF-TEST-1.docx",
        ])

        r = self.client.post(reverse("grading:results-email-preview"), {"audience": "supervisors"}, format="json")
        self.assertEqual(r.json()["to"], "Sam Lee")
        self.assertIn("Dear Sam Lee,", r.json()["html"])
        self.assertEqual(r.json()["attachments"][-1], f"{year}_BTF_Student_Marks.xlsx")
        self.assertEqual(mail.outbox, [])

    def test_preview_names_example_files_while_nobody_is_due(self):
        Submission.objects.update(submitted_at=None)
        r = self.client.post(reverse("grading:results-email-preview"), {"audience": "groups"}, format="json")
        year = self.group.year
        self.assertEqual(r.json()["attachments"], [
            f"{year}_BTF_Certificate_Student_name.docx",
            f"{year}_BTF_Mentor_Certificate_Mentor_name.docx",
            f"{year}_BTF_Marks_Team_name.docx",
        ])
