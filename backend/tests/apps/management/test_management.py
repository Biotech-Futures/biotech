"""Management-tab endpoints: the submission deadline, per-group extensions,
the marks/certificates release gates, and grading settings (director names,
docx templates, signatures)."""
import base64
import io
import zipfile
from datetime import timedelta

from django.core.files.storage import default_storage
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from django.utils import timezone
from rest_framework import status
from rest_framework.test import APIClient

from apps.grading.models import GradingSettings
from apps.submissions.models import Deadline, GroupExtension

from tests.apps.grading.fixtures import _GradingFixture, _seed_doc_templates


class SubmissionDeadlineViewTests(_GradingFixture):
    def setUp(self):
        self.client = APIClient()
        self.url = reverse("grading:deadline")

    def test_non_staff_denied(self):
        self.client.force_authenticate(self.non_staff)
        self.assertEqual(self.client.get(self.url).status_code, status.HTTP_403_FORBIDDEN)

    def test_get_null_when_unset(self):
        self.client.force_authenticate(self.staff)
        r = self.client.get(self.url)
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        self.assertIsNone(r.json()["deadline"])

    def test_post_creates_active_deadline(self):
        self.client.force_authenticate(self.staff)
        r = self.client.post(
            self.url,
            {"closes_at": "2026-10-30T13:00:00Z", "grace_hours": 6},
            format="json",
        )
        self.assertEqual(r.status_code, status.HTTP_200_OK, r.content)
        deadline = r.json()["deadline"]
        self.assertTrue(deadline["is_open"])
        self.assertEqual(deadline["grace_hours"], 6)

        # Newest active row wins: setting again replaces what's in force.
        r2 = self.client.post(
            self.url, {"closes_at": "2020-01-01T00:00:00Z"}, format="json"
        )
        self.assertFalse(r2.json()["deadline"]["is_open"])

    def test_bad_payload_rejected(self):
        self.client.force_authenticate(self.staff)
        self.assertEqual(
            self.client.post(self.url, {"closes_at": "not-a-date"}, format="json").status_code,
            status.HTTP_400_BAD_REQUEST,
        )
        self.assertEqual(
            self.client.post(
                self.url, {"closes_at": "2026-10-30T13:00:00Z", "grace_hours": -1},
                format="json",
            ).status_code,
            status.HTTP_400_BAD_REQUEST,
        )


class GroupExtensionViewTests(_GradingFixture):
    def setUp(self):
        self.client = APIClient()
        self.url = reverse("grading:deadline-extensions")

    def test_non_staff_denied(self):
        self.client.force_authenticate(self.non_staff)
        self.assertEqual(self.client.get(self.url).status_code, status.HTTP_403_FORBIDDEN)

    def test_grant_list_and_revoke(self):
        self.client.force_authenticate(self.staff)
        r = self.client.post(
            self.url,
            {"group_id": self.group.id, "extended_until": "2026-11-05T13:00:00Z",
             "reason": "School flood."},
            format="json",
        )
        self.assertEqual(r.status_code, status.HTTP_200_OK, r.content)
        self.assertEqual(r.json()["extension"]["group_name"], "BTF-TEST-1")

        listed = self.client.get(self.url).json()["extensions"]
        self.assertEqual(len(listed), 1)
        self.assertEqual(listed[0]["reason"], "School flood.")
        self.assertEqual(listed[0]["granted_by"], "Ada Grader")

        # Granting again revokes the old extension and adds a fresh row —
        # the active one leads, the revoked one trails as history.
        self.client.post(
            self.url,
            {"group_id": self.group.id, "extended_until": "2026-11-08T13:00:00Z"},
            format="json",
        )
        listed = self.client.get(self.url).json()["extensions"]
        self.assertEqual(len(listed), 2)
        self.assertIsNone(listed[0]["revoked_at"])
        self.assertIn("2026-11-08", listed[0]["extended_until"])
        self.assertEqual(listed[1]["revoked_by"], "Ada Grader")
        self.assertIn("2026-11-05", listed[1]["extended_until"])

        detail = reverse("grading:deadline-extension-detail", kwargs={"group_id": self.group.id})
        self.assertEqual(self.client.delete(detail).status_code, status.HTTP_204_NO_CONTENT)
        # Idempotent revoke.
        self.assertEqual(self.client.delete(detail).status_code, status.HTTP_204_NO_CONTENT)
        # Soft revoke: both rows stay listed (revoked), and the team's
        # window is no longer extended.
        listed = self.client.get(self.url).json()["extensions"]
        self.assertEqual(len(listed), 2)
        self.assertTrue(all(row["revoked_at"] is not None for row in listed))
        from apps.submissions.services import deadline_for_group
        self.assertFalse(deadline_for_group(self.group.id).is_extended)

        # Re-granting creates a fresh active row beside the history.
        self.client.post(
            self.url,
            {"group_id": self.group.id, "extended_until": "2026-11-09T13:00:00Z"},
            format="json",
        )
        listed = self.client.get(self.url).json()["extensions"]
        self.assertEqual(len(listed), 3)
        self.assertIsNone(listed[0]["revoked_at"])
        self.assertIn("2026-11-09", listed[0]["extended_until"])
        self.assertTrue(all(row["revoked_at"] is not None for row in listed[1:]))
        self.assertTrue(deadline_for_group(self.group.id).is_extended)

    def test_unknown_group_404(self):
        self.client.force_authenticate(self.staff)
        r = self.client.post(
            self.url,
            {"group_id": 999_999, "extended_until": "2026-11-05T13:00:00Z"},
            format="json",
        )
        self.assertEqual(r.status_code, status.HTTP_404_NOT_FOUND)

    def test_extension_before_deadline_rejected(self):
        from datetime import datetime, timezone as tz

        from apps.submissions.models import Deadline

        Deadline.objects.create(
            closes_at=datetime(2026, 10, 1, 13, 0, tzinfo=tz.utc), is_active=True,
        )
        self.client.force_authenticate(self.staff)

        # Earlier than the deadline would shorten the window — refused.
        r = self.client.post(
            self.url,
            {"group_id": self.group.id, "extended_until": "2026-09-15T13:00:00Z"},
            format="json",
        )
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("later than the current deadline", r.json()["detail"])

        # Later is fine.
        r = self.client.post(
            self.url,
            {"group_id": self.group.id, "extended_until": "2026-10-15T13:00:00Z"},
            format="json",
        )
        self.assertEqual(r.status_code, status.HTTP_200_OK, r.content)


class MarksReleaseViewTests(_GradingFixture):
    def setUp(self):
        self.client = APIClient()

    def test_non_staff_denied(self):
        self.client.force_authenticate(self.non_staff)
        resp = self.client.get(reverse("grading:release"))
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)

    def test_flip_on_then_off(self):
        self.client.force_authenticate(self.staff)
        r = self.client.post(reverse("grading:release"), {}, format="json")
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        self.assertIsNotNone(r.json()["released_at"])

        r = self.client.post(reverse("grading:release"), {"release": "false"}, format="json")
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        self.assertIsNone(r.json()["released_at"])

    def test_release_blocked_while_submissions_open(self):
        """Neither gate can be flipped on before the submission window closes."""
        self.client.force_authenticate(self.staff)
        Deadline.objects.create(
            closes_at=timezone.now() + timedelta(days=1), grace_hours=0, is_active=True
        )
        for name in ("grading:release", "grading:certificates-release"):
            r = self.client.post(reverse(name), {}, format="json")
            self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST, name)
            current = self.client.get(reverse(name)).json()
            self.assertIsNone(current["released_at"])
            # The pages use this flag to disable Release before any dialog.
            self.assertTrue(current["submissions_open"])

    def test_release_blocked_while_an_extension_is_open(self):
        """A single team's extension keeps the gates shut past the baseline."""
        self.client.force_authenticate(self.staff)
        Deadline.objects.create(
            closes_at=timezone.now() - timedelta(days=1), grace_hours=0, is_active=True
        )
        GroupExtension.objects.create(
            group=self.group, extended_until=timezone.now() + timedelta(hours=2)
        )
        r = self.client.post(reverse("grading:release"), {}, format="json")
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)

    def test_release_allowed_once_window_has_closed(self):
        self.client.force_authenticate(self.staff)
        Deadline.objects.create(
            closes_at=timezone.now() - timedelta(days=2), grace_hours=24, is_active=True
        )
        GroupExtension.objects.create(
            group=self.group, extended_until=timezone.now() - timedelta(hours=1)
        )
        r = self.client.post(reverse("grading:release"), {}, format="json")
        self.assertEqual(r.status_code, status.HTTP_200_OK, r.content)
        self.assertIsNotNone(r.json()["released_at"])
        self.assertFalse(r.json()["submissions_open"])

    def test_unrelease_allowed_even_while_open(self):
        """Turning a gate OFF is an emergency redaction — never blocked."""
        self.client.force_authenticate(self.staff)
        Deadline.objects.create(
            closes_at=timezone.now() + timedelta(days=1), grace_hours=0, is_active=True
        )
        r = self.client.post(reverse("grading:release"), {"release": "false"}, format="json")
        self.assertEqual(r.status_code, status.HTTP_200_OK)


class GradingSettingsViewTests(_GradingFixture):
    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.staff)

    def test_get_and_patch_director_names(self):
        r = self.client.patch(
            reverse("grading:settings"),
            {"director_1_name": "Alice A", "director_2_name": "Bob B"},
            format="json",
        )
        self.assertEqual(r.status_code, status.HTTP_200_OK, r.content)
        self.assertEqual(r.json()["director_1_name"], "Alice A")

        r2 = self.client.get(reverse("grading:settings"))
        self.assertEqual(r2.json()["director_2_name"], "Bob B")

    def test_template_test_render_streams_synthetic_docx(self):
        _seed_doc_templates()
        for kind in ("marks-summary", "certificate"):
            r = self.client.get(
                reverse("grading:settings-test-render", kwargs={"kind": kind})
            )
            self.assertEqual(r.status_code, status.HTTP_200_OK, kind)
            self.assertIn("wordprocessingml", r["Content-Type"])
            self.assertTrue(r.content[:4] == b"PK\x03\x04", kind)
            # The synthetic values must actually land in the document.
            xml = zipfile.ZipFile(io.BytesIO(r.content)).read("word/document.xml")
            self.assertIn(b"SAMPLE-TEAM-01" if kind == "marks-summary" else b"Jane", xml)

    def test_template_test_render_404s_when_nothing_uploaded(self):
        r = self.client.get(
            reverse("grading:settings-test-render", kwargs={"kind": "marks-summary"})
        )
        self.assertEqual(r.status_code, status.HTTP_404_NOT_FOUND, r.content)

    def test_template_scan_empty_when_nothing_uploaded(self):
        for kind in ("marks-summary", "certificate"):
            r = self.client.get(
                reverse("grading:settings-template-scan", kwargs={"kind": kind})
            )
            self.assertEqual(r.status_code, status.HTTP_200_OK, r.content)
            self.assertEqual(
                r.json(),
                {"uploaded": False, "dialect": "none", "present": [], "unknown": []},
            )

    def test_template_scan_flags_unrecognised_tokens(self):
        from docx import Document as NewDocument

        from apps.grading.models import GradingSettings

        doc = NewDocument()
        doc.add_paragraph("{{TeamCode}} {{Typoed}} {{Director1Signature}}")
        buf = io.BytesIO()
        doc.save(buf)

        settings_row = GradingSettings.load()
        settings_row.marks_summary_template = SimpleUploadedFile("tpl.docx", buf.getvalue())
        settings_row.save()

        data = self.client.get(
            reverse("grading:settings-template-scan", kwargs={"kind": "marks-summary"})
        ).json()
        self.assertTrue(data["uploaded"])
        self.assertIn("TeamCode", data["present"])
        self.assertIn("Director1Signature", data["present"])
        self.assertEqual(data["unknown"], ["Typoed"])
        # Placeholders the template does not use are absent, not reported.
        self.assertNotIn("SAQTotal", data["present"])

    def test_template_scan_unknown_kind_404s(self):
        r = self.client.get(
            reverse("grading:settings-template-scan", kwargs={"kind": "poster"})
        )
        self.assertEqual(r.status_code, status.HTTP_404_NOT_FOUND)

    def test_template_test_render_unknown_kind_404s(self):
        r = self.client.get(
            reverse("grading:settings-test-render", kwargs={"kind": "poster"})
        )
        self.assertEqual(r.status_code, status.HTTP_404_NOT_FOUND)

    def test_template_test_render_requires_grader(self):
        self.client.force_authenticate(self.non_staff)
        r = self.client.get(
            reverse("grading:settings-test-render", kwargs={"kind": "certificate"})
        )
        self.assertEqual(r.status_code, status.HTTP_403_FORBIDDEN)

    @staticmethod
    def _docx_bytes(text: str) -> bytes:
        from docx import Document as NewDocument

        doc = NewDocument()
        doc.add_paragraph(text)
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    # 1x1 transparent PNG — the smallest thing python-docx accepts as an image.
    _PNG = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAACklEQVR4nGNgAAAA"
        "AgABc3UBGAAAAABJRU5ErkJggg=="
    )

    def test_upload_rejects_file_that_is_not_a_docx(self):
        r = self.client.patch(
            reverse("grading:settings"),
            {"marks_summary_template": SimpleUploadedFile("tpl.docx", b"not a zip at all")},
            format="multipart",
        )
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST, r.content)
        self.assertIn("marks_summary_template", r.json()["fields"])
        self.assertFalse(GradingSettings.load().marks_summary_template)

    def test_upload_rejects_signature_that_is_not_an_image(self):
        r = self.client.patch(
            reverse("grading:settings"),
            {"director_1_signature": SimpleUploadedFile("sig.png", b"plainly not pixels")},
            format="multipart",
        )
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST, r.content)
        self.assertIn("director_1_signature", r.json()["fields"])
        self.assertFalse(GradingSettings.load().director_1_signature)

    def test_bad_file_rejects_the_whole_patch(self):
        """One bad upload must not let the rest of the request half-apply."""
        r = self.client.patch(
            reverse("grading:settings"),
            {
                "director_1_name": "Dr. Half Applied",
                "certificate_template": SimpleUploadedFile("c.docx", b"broken"),
            },
            format="multipart",
        )
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST, r.content)
        self.assertEqual(GradingSettings.load().director_1_name, "")

    def test_replacing_template_deletes_the_old_blob(self):
        first = self.client.patch(
            reverse("grading:settings"),
            {
                "marks_summary_template": SimpleUploadedFile(
                    "a.docx", self._docx_bytes("{{TeamCode}}")
                )
            },
            format="multipart",
        )
        self.assertEqual(first.status_code, status.HTTP_200_OK, first.content)
        old_name = GradingSettings.load().marks_summary_template.name
        self.assertTrue(default_storage.exists(old_name))

        second = self.client.patch(
            reverse("grading:settings"),
            {
                "marks_summary_template": SimpleUploadedFile(
                    "b.docx", self._docx_bytes("{{TeamCode}}")
                )
            },
            format="multipart",
        )
        self.assertEqual(second.status_code, status.HTTP_200_OK, second.content)
        new_name = GradingSettings.load().marks_summary_template.name
        self.assertNotEqual(old_name, new_name)
        self.assertTrue(default_storage.exists(new_name))
        # The superseded blob is cleaned up, not orphaned in the container.
        self.assertFalse(default_storage.exists(old_name))

    def test_clearing_a_signature_deletes_its_blob(self):
        upload = self.client.patch(
            reverse("grading:settings"),
            {"director_1_signature": SimpleUploadedFile("sig.png", self._PNG)},
            format="multipart",
        )
        self.assertEqual(upload.status_code, status.HTTP_200_OK, upload.content)
        name = GradingSettings.load().director_1_signature.name
        self.assertTrue(default_storage.exists(name))

        cleared = self.client.patch(
            reverse("grading:settings"), {"director_1_signature": None}, format="json"
        )
        self.assertEqual(cleared.status_code, status.HTTP_200_OK, cleared.content)
        self.assertFalse(GradingSettings.load().director_1_signature)
        self.assertFalse(default_storage.exists(name))

    def test_candidate_scan_previews_without_saving(self):
        r = self.client.post(
            reverse("grading:settings-template-scan", kwargs={"kind": "marks-summary"}),
            {
                "file": SimpleUploadedFile(
                    "draft.docx", self._docx_bytes("{{TeamCode}} {{Typoed}}")
                )
            },
            format="multipart",
        )
        self.assertEqual(r.status_code, status.HTTP_200_OK, r.content)
        data = r.json()
        self.assertEqual(data["dialect"], "tokens")
        self.assertIn("TeamCode", data["present"])
        self.assertEqual(data["unknown"], ["Typoed"])
        # Preview only — nothing was stored.
        self.assertFalse(GradingSettings.load().marks_summary_template)

    def test_candidate_scan_rejects_unreadable_file(self):
        r = self.client.post(
            reverse("grading:settings-template-scan", kwargs={"kind": "certificate"}),
            {"file": SimpleUploadedFile("draft.docx", b"broken")},
            format="multipart",
        )
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST, r.content)
        self.assertFalse(GradingSettings.load().certificate_template)

    def test_candidate_test_render_streams_without_saving(self):
        r = self.client.post(
            reverse("grading:settings-test-render", kwargs={"kind": "marks-summary"}),
            {
                "file": SimpleUploadedFile(
                    "draft.docx", self._docx_bytes("Team {{TeamCode}}")
                )
            },
            format="multipart",
        )
        self.assertEqual(r.status_code, status.HTTP_200_OK, r.content)
        self.assertIn("wordprocessingml", r["Content-Type"])
        # The synthetic team code must land in the candidate's rendered copy.
        xml = zipfile.ZipFile(io.BytesIO(r.content)).read("word/document.xml")
        self.assertIn(b"SAMPLE-TEAM-01", xml)
        self.assertFalse(GradingSettings.load().marks_summary_template)

    def test_curly_tokens_scan_and_render(self):
        """The {{Name}} token syntax works through scan and test render."""
        scan = self.client.post(
            reverse("grading:settings-template-scan", kwargs={"kind": "marks-summary"}),
            {
                "file": SimpleUploadedFile(
                    "draft.docx", self._docx_bytes("{{TeamCode}} {{Typoed}}")
                )
            },
            format="multipart",
        )
        self.assertEqual(scan.status_code, status.HTTP_200_OK, scan.content)
        self.assertEqual(scan.json()["dialect"], "tokens")
        self.assertIn("TeamCode", scan.json()["present"])
        self.assertEqual(scan.json()["unknown"], ["Typoed"])

        render = self.client.post(
            reverse("grading:settings-test-render", kwargs={"kind": "marks-summary"}),
            {"file": SimpleUploadedFile("draft.docx", self._docx_bytes("Team {{TeamCode}}"))},
            format="multipart",
        )
        self.assertEqual(render.status_code, status.HTTP_200_OK, render.content)
        xml = zipfile.ZipFile(io.BytesIO(render.content)).read("word/document.xml")
        self.assertIn(b"SAMPLE-TEAM-01", xml)
        self.assertNotIn(b"{{TeamCode}}", xml)

    def test_candidate_post_without_file_400s(self):
        r = self.client.post(
            reverse("grading:settings-template-scan", kwargs={"kind": "marks-summary"}),
            {},
            format="multipart",
        )
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)
