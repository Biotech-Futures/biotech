"""Results emails from the Release Results tab: who gets which email, the
client's wording, the files they carry, the survey details, preview, and
sending each audience in batches."""
import io
import zipfile
from datetime import timedelta
from decimal import Decimal
from email.mime.base import MIMEBase
from unittest import mock

from django.core import mail
from django.test import override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from rest_framework import status
from rest_framework.test import APIClient

from apps.grading.models import FinalistFlag, Grade, Rubric, RubricCriterion, SubmissionComponent
from apps.groups.models import GroupMembership, Groups
from apps.management.models import (
    CertificatesRelease,
    GradingSettings,
    MarksRelease,
    ResultsEmailSettings,
    ResultsSupervisorEmail,
    ResultsTeamEmail,
)
from apps.management.services.finalist_notify import symposium_today
from apps.services.models import SystemEmailTemplate
from apps.submissions.models import Submission
from apps.users.models import StudentProfile, SupervisorProfile, User

from tests.apps.grading.fixtures import _GradingFixture
from tests.apps.management.fixtures import _seed_doc_templates

LOCMEM = "django.core.mail.backends.locmem.EmailBackend"
SEND = "management:results-email-send"


def _user(email, first="", last=""):
    return User.objects.create_user(email=email, first_name=first, last_name=last, password="pw12345!")


def _student(email, group, supervisor_profile=None):
    user = _user(email, "Stu", email.split("@")[0])
    StudentProfile.objects.create(
        user=user, pg_first_name="P", pg_last_name="G", parent_guardian_flag=True,
        school_name="Test School", year_lvl="10", supervisor=supervisor_profile,
    )
    GroupMembership.objects.create(user=user, group=group, membership_role="student")
    return user


def _files(message):
    """A sent email's files by name (the inline logo left out)."""
    return {name: content for name, content, _ in (a for a in message.attachments if not isinstance(a, MIMEBase))}


def _docx_text(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return z.read("word/document.xml").decode("utf8")


def _submitted_team(name, staff):
    team = Groups.objects.create(group_name=name)
    submission = Submission.objects.create(group=team, answers={"q_answers": "Answers."})
    submission.snapshot(staff)
    submission.save()
    return team


@override_settings(EMAIL_BACKEND=LOCMEM)
class ResultsEmailTests(_GradingFixture):
    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.staff)

        # The fixture's submitted group: two students, a mentor, and a supervisor
        # who is also a member. The students and the mentor get the group email.
        self.supervisor = _user("sam.lee@example.com", "Sam", "Lee")
        self.sup_profile = SupervisorProfile.objects.create(user=self.supervisor)
        _student("amy@example.com", self.group, self.sup_profile)
        _student("ben@example.com", self.group, self.sup_profile)
        mentor = _user("mentor@example.com", "Mo", "Mentor")
        GroupMembership.objects.create(user=mentor, group=self.group, membership_role="mentor")
        GroupMembership.objects.create(user=self.supervisor, group=self.group, membership_role="supervisor")

        # A finalist team (left out while certificates go out without finalists)
        # and a team that never submitted.
        self.finalist = _submitted_team("Finalist Team", self.staff)
        _student("fin@example.com", self.finalist, self.sup_profile)
        FinalistFlag.objects.create(group=self.finalist, flagged_by=self.staff)
        no_entry = Groups.objects.create(group_name="No Entry")
        _student("none@example.com", no_entry)

        now = timezone.now()
        marks = MarksRelease.load()
        marks.released_at = now
        marks.save()
        certificates = CertificatesRelease.load()
        certificates.released_at = now
        certificates.exclude_finalists = True
        certificates.save()
        details = ResultsEmailSettings.load()
        details.survey_url = "https://example.com/survey"
        details.survey_closes = symposium_today() + timedelta(days=30)
        details.save()
        _seed_doc_templates()
        Grade.objects.create(
            submission=self.submission, criterion=self.saq_c1, mark=Decimal("8.00"), graded_by=self.staff,
        )
        Grade.objects.create(
            submission=self.submission, criterion=self.poster_c1, mark=Decimal("6.50"), graded_by=self.staff,
        )

    def _send_all(self, audience):
        """One press: the run sends everyone (inline under test settings). The
        run's result, with the counts after it, as a one-item list."""
        r = self.client.post(reverse(SEND), {"audience": audience}, format="json")
        self.assertEqual(r.status_code, status.HTTP_200_OK, r.content)
        body = r.json()
        run = body["runs"][audience]
        self.assertFalse(run["sending"])
        return [{**body, "emailed": run["run"]["emailed"], "failed": run["run"]["failed"], "run": run["run"]}]

    def _recipients(self):
        return sorted(m.to[0] for m in mail.outbox)

    # -- who gets it ------------------------------------------------------------

    def test_the_groups_students_and_mentors_get_the_group_email(self):
        results = self._send_all("groups")
        # The supervisor who is a member gets the supervisor email instead.
        self.assertEqual(self._recipients(), ["amy@example.com", "ben@example.com", "mentor@example.com"])

        message = mail.outbox[0]
        text = " ".join(message.body.split())  # the plain text wraps its lines
        self.assertEqual(message.subject, f"Your {self.group.year} BIOTech Futures Challenge results")
        self.assertIn(f"Dear {self.group.group_name},", text)
        self.assertIn(f"Congratulations on your participation in the {self.group.year} BIOTech Futures Challenge!", text)
        self.assertIn("https://example.com/survey", text)
        closes = ResultsEmailSettings.load().survey_closes
        self.assertIn(f"until the {closes.day}", text)
        self.assertIn(f" of {closes:%B}.", text)
        self.assertEqual(message.reply_to, ["support@biotechfutures.org"])

        self.assertEqual(results[-1]["emailed"], 3)
        self.assertEqual(results[-1]["groups"], {"total": 1, "emailed": 1})
        self.assertEqual(results[-1]["supervisors"], {"total": 1, "emailed": 0})
        self.assertTrue(ResultsTeamEmail.objects.filter(group=self.group).exists())

    def test_supervisors_are_emailed_separately_once_each(self):
        results = self._send_all("supervisors")
        self.assertEqual(self._recipients(), ["sam.lee@example.com"])
        message = mail.outbox[0]
        self.assertEqual(message.subject, f"Your students’ {self.group.year} BIOTech Futures Challenge results")
        self.assertIn("Dear Sam Lee,", message.body)
        self.assertIn("Merit certificates for each of your students", " ".join(message.body.split()))
        self.assertEqual(message.reply_to, ["support@biotechfutures.org"])
        self.assertEqual(results[-1]["supervisors"], {"total": 1, "emailed": 1})
        self.assertEqual(results[-1]["groups"], {"total": 1, "emailed": 0})
        self.assertTrue(ResultsSupervisorEmail.objects.filter(supervisor=self.supervisor).exists())

    def test_finalists_are_included_when_certificates_go_to_everyone(self):
        CertificatesRelease.objects.update(exclude_finalists=False)
        self._send_all("groups")
        self._send_all("supervisors")
        self.assertIn("fin@example.com", self._recipients())
        # Their supervisor is the same person: still one supervisor email.
        self.assertEqual(self._recipients().count("sam.lee@example.com"), 1)

    def test_a_second_send_emails_nobody_again(self):
        self._send_all("groups")
        self._send_all("supervisors")
        mail.outbox = []
        self.assertEqual(self._send_all("groups")[-1]["run"]["due"], 0)
        self.assertEqual(self._send_all("supervisors")[-1]["run"]["due"], 0)
        self.assertEqual(mail.outbox, [])

    # -- one run at a time, on the server ------------------------------------------

    def _hold_send(self, key, until):
        from apps.management.models import EmailSendRun

        EmailSendRun.objects.update_or_create(key=key, defaults={"held_until": until, "started_at": timezone.now()})

    def _runs(self):
        return self.client.get(reverse("management:results-email")).json()["runs"]

    def test_a_second_run_is_refused_while_one_is_sending(self):
        # Another tab, or a run started before the page was reopened.
        self._hold_send("results_team", timezone.now() + timedelta(minutes=4))
        self.assertEqual((self._runs()["groups"]["sending"], self._runs()["supervisors"]["sending"]), (True, False))
        r = self.client.post(reverse(SEND), {"audience": "groups"}, format="json")
        self.assertEqual(r.status_code, status.HTTP_409_CONFLICT)
        self.assertEqual(r.json()["detail"], "Results emails are already being sent. Wait for that to finish.")
        self.assertEqual(mail.outbox, [])
        self.assertFalse(ResultsTeamEmail.objects.exists())
        # The supervisor email has a run of its own.
        self._send_all("supervisors")
        self.assertEqual(self._recipients(), ["sam.lee@example.com"])

    def test_the_page_is_told_the_runs_progress(self):
        self.assertEqual(self._runs()["groups"], {"sending": False, "run": None})
        self._send_all("groups")
        run = self._runs()["groups"]
        self.assertFalse(run["sending"])
        # Two students and a mentor due, all emailed.
        self.assertEqual(
            {k: run["run"][k] for k in ("due", "emailed", "failed", "error")},
            {"due": 3, "emailed": 3, "failed": 0, "error": ""},
        )
        self.assertIsNotNone(run["run"]["finished_at"])

    def test_a_run_that_died_frees_it_once_its_hold_lapses(self):
        self._hold_send("results_team", timezone.now() - timedelta(seconds=1))
        self.assertFalse(self._runs()["groups"]["sending"])
        self._send_all("groups")
        self.assertEqual(self._recipients(), ["amy@example.com", "ben@example.com", "mentor@example.com"])

    def test_an_extension_granted_since_the_release_holds_them(self):
        from apps.submissions.models import GroupExtension

        GroupExtension.objects.create(group=self.group, extended_until=timezone.now() + timedelta(hours=6))
        body = self.client.get(reverse("management:results-email")).json()
        self.assertTrue(body["submissions_open"].startswith("A team's extension is open until "), body)
        for audience in ("groups", "supervisors"):
            r = self.client.post(reverse(SEND), {"audience": audience}, format="json")
            self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)
            self.assertEqual(r.json()["detail"], body["submissions_open"])
        self.assertEqual(mail.outbox, [])

    def test_a_run_that_cannot_reach_the_mail_server_says_so_and_frees_the_send(self):
        unreachable = mock.Mock()
        unreachable.open.side_effect = OSError("mail server down")
        with mock.patch("apps.management.services.send_guard.get_connection", return_value=unreachable), \
                self.assertLogs("apps.management.services.send_guard", level="ERROR"):
            run = self._send_all("groups")[-1]["run"]
        self.assertEqual(
            (run["emailed"], run["failed"], run["error"]),
            (0, 1, "Couldn't reach the mail server. Press Send again to email the rest."),
        )
        # Nobody it was for got it, so all are listed.
        group = self.group.group_name
        self.assertEqual(
            sorted(run["missed"]), [f"({group}) Mo Mentor", f"({group}) Stu amy", f"({group}) Stu ben"],
        )
        self.assertFalse(ResultsTeamEmail.objects.exists())
        # Pressing again sends it.
        self._send_all("groups")
        self.assertTrue(ResultsTeamEmail.objects.filter(group=self.group).exists())

    # -- the files they carry ------------------------------------------------------

    def test_everyone_in_the_group_gets_all_its_certificates_and_the_marks_summary(self):
        self._send_all("groups")
        year = self.group.year
        self.assertEqual(len(mail.outbox), 3)
        for message in mail.outbox:
            # Students and mentor alike: each sees the others' certificates.
            files = _files(message)
            self.assertEqual(list(files), [
                f"{year}_BTF_Student_Certificate_Stu_amy.docx",
                f"{year}_BTF_Student_Certificate_Stu_ben.docx",
                f"{year}_BTF_Mentor_Certificate_Mo_Mentor.docx",
                f"{year}_BTF_Marks_BTF-TEST-1.docx",
            ])
            self.assertIn("Stu amy — BTF-TEST-1", _docx_text(files[f"{year}_BTF_Student_Certificate_Stu_amy.docx"]))
            self.assertIn(
                "Mo Mentor mentored BTF-TEST-1", _docx_text(files[f"{year}_BTF_Mentor_Certificate_Mo_Mentor.docx"]),
            )
            self.assertIn("Team BTF-TEST-1 S1 8 ", _docx_text(files[f"{year}_BTF_Marks_BTF-TEST-1.docx"]))

    def test_the_files_come_after_the_email_and_its_logo(self):
        self._send_all("groups")
        top = mail.outbox[0].message()
        self.assertEqual(top.get_content_type(), "multipart/mixed")
        body, *files = top.get_payload()
        self.assertEqual(body.get_content_type(), "multipart/related")
        self.assertEqual(
            [part.get_content_type() for part in body.get_payload()],
            ["multipart/alternative", "image/png"],
        )
        self.assertEqual([f.get_content_disposition() for f in files], ["attachment"] * 4)

    def test_a_comment_pasted_from_word_still_reaches_the_supervisor(self):
        # A vertical tab (Word's Shift+Enter) and a stray control character.
        Grade.objects.filter(criterion=self.poster_c1).update(comment="Bold\x0blayout.\x01")
        run = self._send_all("supervisors")[-1]["run"]
        self.assertEqual((run["emailed"], run["failed"], run["missed"]), (1, 0, []))
        files = _files(mail.outbox[0])
        sheet = load_workbook(io.BytesIO(files[f"{self.group.year}_BTF_Student_Marks_Sam_Lee.xlsx"])).active
        rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
        self.assertEqual(dict(zip(rows[0], rows[1]))["PosterComment1"], "Bold layout.")

    def test_supervisors_get_their_students_and_mentors_certificates_and_a_marks_sheet(self):
        # Only Poster and SAQ, as on the marks summary: another marked component stays out.
        report = SubmissionComponent.objects.exclude(code__in=("SAQ", "POSTER")).first()
        RubricCriterion.objects.create(
            rubric=Rubric.objects.create(component=report, year=self.year, active=True),
            name="Structure", max_mark=Decimal("5.00"), order=10,
        )
        Grade.objects.filter(criterion=self.poster_c1).update(comment="Bold, clear layout.")
        self._send_all("supervisors")
        year = self.group.year
        files = _files(mail.outbox[0])
        # Their finalist student is left out while certificates exclude finalists.
        self.assertEqual(list(files), [
            f"{year}_BTF_Student_Certificate_Stu_amy.docx",
            f"{year}_BTF_Student_Certificate_Stu_ben.docx",
            f"{year}_BTF_Mentor_Certificate_Mo_Mentor.docx",
            f"{year}_BTF_Student_Marks_Sam_Lee.xlsx",
        ])
        self.assertIn(
            "Mo Mentor mentored BTF-TEST-1", _docx_text(files[f"{year}_BTF_Mentor_Certificate_Mo_Mentor.docx"]),
        )

        # One row per group, in the marks summary's own field names: each mark
        # followed by its comment, then the total.
        sheet = load_workbook(io.BytesIO(files[f"{year}_BTF_Student_Marks_Sam_Lee.xlsx"])).active
        rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
        self.assertEqual(rows[0], [
            "TeamCode", "Students", "Mentor", "ProjectTitle", "ProjectCategory", "SolutionCategory",
            *(name for i in range(1, 11) for name in (f"PM{i}", f"PosterComment{i}")),
            *(name for i in range(1, 5) for name in (f"SM{i}", f"ShortAnswerQuestionComment{i}")),
            "MTotal",
        ])
        self.assertEqual(rows[0][6:10], ["PM1", "PosterComment1", "PM2", "PosterComment2"])
        self.assertEqual(len(rows), 2)
        row = dict(zip(rows[0], rows[1]))
        self.assertEqual(
            (row["TeamCode"], row["Students"], row["Mentor"]),
            ("BTF-TEST-1", "Stu amy, Stu ben", "Mo Mentor"),
        )
        # The design mark is poster criterion 1; content is SAQ 1, clarity (SAQ 2) unmarked.
        self.assertEqual((row["PM1"], row["PM2"], row["SM1"], row["SM2"]), (6.5, None, 8, None))
        # A comment sits beside its mark; one not given stays blank.
        self.assertEqual(
            (row["PosterComment1"], row["ShortAnswerQuestionComment1"]), ("Bold, clear layout.", None)
        )
        # Comments are wide and wrap.
        comment_letter = get_column_letter(rows[0].index("PosterComment1") + 1)
        self.assertEqual(sheet.column_dimensions[comment_letter].width, 30)
        self.assertTrue(sheet[f"{comment_letter}2"].alignment.wrap_text)
        # One total: the Poster and SAQ marks together.
        self.assertEqual(row["MTotal"], 14.5)

    def test_a_team_whose_files_cant_be_made_stays_pending(self):
        with mock.patch(
            "apps.management.services.results_notify.Documents.marks_summary", side_effect=OSError("storage down"),
        ), self.assertLogs("apps.management.services.results_notify", level="ERROR"):
            results = self._send_all("groups")
        self.assertEqual(mail.outbox, [])
        self.assertEqual(results[-1]["failed"], 1)
        self.assertFalse(ResultsTeamEmail.objects.filter(group=self.group).exists())

    def test_sending_waits_for_the_templates_its_files_need(self):
        GradingSettings.objects.update(marks_summary_template="")
        r = self.client.post(reverse(SEND), {"audience": "groups"}, format="json")
        self.assertEqual(
            r.json()["detail"],
            "Upload the marks summary, student certificate and mentor certificate templates "
            "in Document Setup before emailing groups.",
        )
        # The supervisors' spreadsheet needs no template.
        self._send_all("supervisors")
        self.assertEqual(self._recipients(), ["sam.lee@example.com"])

        GradingSettings.objects.update(mentor_certificate_template="")
        r = self.client.post(reverse(SEND), {"audience": "supervisors"}, format="json")
        self.assertEqual(
            r.json()["detail"],
            "Upload the student certificate and mentor certificate templates "
            "in Document Setup before emailing supervisors.",
        )
        r = self.client.get(reverse("management:results-email"))
        self.assertEqual(r.json()["templates_ready"], {"groups": False, "supervisors": False})

    # -- runs and failures ----------------------------------------------------------

    def test_a_supervisor_it_could_not_reach_is_listed_with_their_groups(self):
        with mock.patch("django.core.mail.EmailMultiAlternatives.send", autospec=True, side_effect=OSError("rejected")), \
                self.assertLogs("apps.management.services.results_notify", level="ERROR"):
            run = self._send_all("supervisors")[-1]["run"]
        self.assertEqual(run["missed"], [f"({self.group.group_name}) Sam Lee"])

    def test_one_press_emails_every_group(self):
        for n in range(6):
            team = _submitted_team(f"Extra {n}", self.staff)
            _student(f"extra{n}@example.com", team)
        result = self._send_all("groups")[-1]
        # The fixture group has two students and a mentor; six more students.
        self.assertEqual((result["run"]["due"], result["emailed"]), (9, 9))
        self.assertEqual(result["groups"], {"total": 7, "emailed": 7})

    def test_a_team_whose_student_misses_it_stays_pending_and_is_not_retried_in_the_run(self):
        real_send = mail.EmailMultiAlternatives.send

        def fail_for_ben(message, *args, **kwargs):
            if message.to == ["ben@example.com"]:
                raise OSError("rejected")
            return real_send(message, *args, **kwargs)

        with mock.patch("django.core.mail.EmailMultiAlternatives.send", autospec=True, side_effect=fail_for_ben), \
                self.assertLogs("apps.management.services.results_notify", level="ERROR"):
            results = self._send_all("groups")
        self.assertEqual(sum(r["failed"] for r in results), 1)
        # The page lists who it couldn't reach.
        self.assertEqual(results[-1]["run"]["missed"], [f"({self.group.group_name}) Stu ben"])
        self.assertFalse(ResultsTeamEmail.objects.filter(group=self.group).exists())
        self.assertEqual(results[-1]["groups"], {"total": 1, "emailed": 0})

        # The next press reaches the team again.
        mail.outbox = []
        self._send_all("groups")
        self.assertTrue(ResultsTeamEmail.objects.filter(group=self.group).exists())

    # -- when it may be sent ----------------------------------------------------------

    def test_sending_waits_for_both_releases(self):
        CertificatesRelease.objects.update(released_at=None)
        for audience in ("groups", "supervisors"):
            r = self.client.post(reverse(SEND), {"audience": audience}, format="json")
            self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)
            self.assertEqual(r.json()["detail"], "Release both marks and certificates before sending the results emails.")
        self.assertEqual(mail.outbox, [])

    def test_only_the_group_email_waits_for_the_survey_details(self):
        ResultsEmailSettings.objects.update(survey_closes=None)
        r = self.client.post(reverse(SEND), {"audience": "groups"}, format="json")
        self.assertEqual(r.json()["detail"], "Set the feedback survey link and close date before emailing groups.")
        ResultsEmailSettings.objects.update(survey_closes=symposium_today() - timedelta(days=1))
        r = self.client.post(reverse(SEND), {"audience": "groups"}, format="json")
        self.assertEqual(r.json()["detail"], "The survey close date can't be before today. Update it before emailing groups.")
        # Supervisors aren't told about the survey, so they can go.
        self._send_all("supervisors")
        self.assertEqual(self._recipients(), ["sam.lee@example.com"])

    def test_the_survey_link_starts_empty_and_holds_the_group_email(self):
        # Last year's survey must never go out unnoticed.
        ResultsEmailSettings.objects.all().delete()
        self.assertEqual(ResultsEmailSettings.load().survey_url, "")
        ResultsEmailSettings.objects.update(survey_closes=symposium_today() + timedelta(days=30))
        r = self.client.post(reverse(SEND), {"audience": "groups"}, format="json")
        self.assertEqual(r.json()["detail"], "Set the feedback survey link and close date before emailing groups.")
        self.assertEqual(mail.outbox, [])

    def test_links_left_from_the_old_defaults_are_cleared(self):
        from importlib import import_module

        from django.apps import apps as registry

        from apps.management.models import FinalistEmailSettings

        class apps:
            """The models as the migration knew them, when they were grading's."""

            @staticmethod
            def get_model(app_label, name):
                return registry.get_model("management", name)

        migration = import_module("apps.grading.migrations.0015_email_links_start_empty")
        ResultsEmailSettings.objects.update(survey_url=migration.OLD_SURVEY_URL)
        FinalistEmailSettings.objects.create(registration_url=migration.OLD_REGISTRATION_URL)
        migration.clear_old_links(apps, None)
        self.assertEqual(ResultsEmailSettings.load().survey_url, "")
        self.assertEqual(FinalistEmailSettings.load().registration_url, "")
        # A link an admin entered stays.
        ResultsEmailSettings.objects.update(survey_url="https://example.com/survey")
        migration.clear_old_links(apps, None)
        self.assertEqual(ResultsEmailSettings.load().survey_url, "https://example.com/survey")

    def test_a_switched_off_email_is_refused(self):
        SystemEmailTemplate.objects.create(key="results_supervisor", is_enabled=False)
        r = self.client.post(reverse(SEND), {"audience": "supervisors"}, format="json")
        self.assertEqual(r.json()["detail"], "Results (To supervisors) is switched off on System Emails.")
        r = self.client.get(reverse("management:results-email"))
        self.assertEqual(r.json()["emails_on"], {"groups": True, "supervisors": False})

    def test_an_unknown_audience_is_refused(self):
        r = self.client.post(reverse(SEND), {"audience": "mentors"}, format="json")
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)

    def test_graders_only(self):
        self.client.force_authenticate(self.non_staff)
        for name in ("management:results-email", SEND):
            self.assertEqual(self.client.post(reverse(name), {}, format="json").status_code, status.HTTP_403_FORBIDDEN)

    # -- details and preview -------------------------------------------------------------

    def test_details_are_saved_and_a_past_close_date_refused(self):
        closes = symposium_today() + timedelta(days=10)
        r = self.client.patch(
            reverse("management:results-email"),
            {"survey_url": "https://example.com/survey", "survey_closes": closes.isoformat()},
            format="json",
        )
        self.assertEqual(r.status_code, status.HTTP_200_OK, r.content)
        body = r.json()
        self.assertEqual((body["survey_url"], body["survey_closes"], body["complete"]), ("https://example.com/survey", closes.isoformat(), True))
        self.assertEqual((body["marks_released"], body["certificates_released"]), (True, True))
        self.assertEqual(body["groups"], {"total": 1, "emailed": 0})
        self.assertEqual(body["supervisors"], {"total": 1, "emailed": 0})

        r = self.client.patch(
            reverse("management:results-email"),
            {"survey_closes": (symposium_today() - timedelta(days=1)).isoformat()},
            format="json",
        )
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)

    def test_preview_shows_each_email_with_unsaved_details(self):
        r = self.client.post(
            reverse("management:results-email-preview"),
            {"audience": "groups", "survey_url": "https://example.com/draft"},
            format="json",
        )
        self.assertEqual(r.status_code, status.HTTP_200_OK, r.content)
        self.assertEqual(r.json()["to"], self.group.group_name)
        self.assertIn('href="https://example.com/draft"', r.json()["html"])
        self.assertEqual(ResultsEmailSettings.load().survey_url, "https://example.com/survey")

        year = self.group.year
        self.assertEqual(r.json()["attachments"], [
            f"{year}_BTF_Student_Certificate_Stu_amy.docx",
            f"{year}_BTF_Student_Certificate_Stu_ben.docx",
            f"{year}_BTF_Mentor_Certificate_Mo_Mentor.docx",
            f"{year}_BTF_Marks_BTF-TEST-1.docx",
        ])

        r = self.client.post(reverse("management:results-email-preview"), {"audience": "supervisors"}, format="json")
        self.assertEqual(r.json()["to"], "Sam Lee")
        self.assertIn("Dear Sam Lee,", r.json()["html"])
        self.assertEqual(r.json()["attachments"][-1], f"{year}_BTF_Student_Marks_Sam_Lee.xlsx")
        self.assertEqual(mail.outbox, [])

    def test_the_sample_spreadsheet_has_made_up_groups_marked_on_the_real_rubric(self):
        r = self.client.get(reverse("management:results-email-sample-sheet"))
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        self.assertEqual(
            r["Content-Disposition"], f'attachment; filename="{self.group.year}_BTF_Student_Marks_Sample.xlsx"',
        )
        rows = [[c.value for c in row] for row in load_workbook(io.BytesIO(r.content)).active.iter_rows()]
        self.assertEqual(rows[0][:6], ["TeamCode", "Students", "Mentor", "ProjectTitle", "ProjectCategory", "SolutionCategory"])
        self.assertEqual(rows[0][-3:], ["SM4", "ShortAnswerQuestionComment4", "MTotal"])
        teams = [dict(zip(rows[0], row)) for row in rows[1:]]
        self.assertEqual(sorted(t["TeamCode"] for t in teams), ["SAMPLE1", "SAMPLE2", "SAMPLE3"])
        self.assertEqual({t["TeamCode"]: t["Mentor"] for t in teams}["SAMPLE1"], "Dr Sample Mentor")
        # Highest total first, as in the real spreadsheet.
        totals = [t["MTotal"] for t in teams]
        self.assertEqual(totals, sorted(totals, reverse=True))
        # The fixture rubric: SAQ criteria out of 10 and 5, one poster criterion out of 10.
        for t in teams:
            self.assertIsNotNone(t["SM1"])
            self.assertIsNotNone(t["SM2"])
            self.assertIsNone(t["SM3"])
            self.assertIsNotNone(t["PM1"])
            self.assertIsNone(t["PM2"])
            self.assertEqual(t["MTotal"], t["PM1"] + t["SM1"] + t["SM2"])
        # Not every sample group scores the same.
        self.assertGreater(len({t["MTotal"] for t in teams}), 1)
        # Nothing about real people.
        self.assertNotIn("Stu amy", str(rows))

    def test_the_files_carry_the_title_the_team_submitted(self):
        Submission.objects.filter(group=self.group).update(submitted_project_title="Plant Sensors")
        self._send_all("groups")
        year = self.group.year
        files = _files(mail.outbox[0])
        self.assertIn("Stu amy — Plant Sensors", _docx_text(files[f"{year}_BTF_Student_Certificate_Stu_amy.docx"]))
        self.assertIn(
            "Mo Mentor mentored Plant Sensors", _docx_text(files[f"{year}_BTF_Mentor_Certificate_Mo_Mentor.docx"])
        )
        mail.outbox.clear()
        self._send_all("supervisors")
        sheet = load_workbook(
            io.BytesIO(_files(mail.outbox[0])[f"{year}_BTF_Student_Marks_Sam_Lee.xlsx"])
        ).active
        rows = [[c.value for c in row] for row in sheet.iter_rows()]
        self.assertEqual(rows[1][rows[0].index("ProjectTitle")], "Plant Sensors")

    def test_the_spreadsheet_puts_the_highest_total_first(self):
        # A second group of Sam's students that scored less, and one not marked.
        lower = _submitted_team("A Lower Team", self.staff)
        _student("lo@example.com", lower, self.sup_profile)
        Grade.objects.create(
            submission=Submission.objects.get(group=lower), criterion=self.saq_c1,
            mark=Decimal("2.00"), graded_by=self.staff,
        )
        unmarked = _submitted_team("An Unmarked Team", self.staff)
        _student("un@example.com", unmarked, self.sup_profile)
        self._send_all("supervisors")
        files = _files(mail.outbox[0])
        sheet = load_workbook(io.BytesIO(files[f"{self.group.year}_BTF_Student_Marks_Sam_Lee.xlsx"])).active
        rows = [[c.value for c in row] for row in sheet.iter_rows()]
        codes = [row[0] for row in rows[1:]]
        totals = [row[-1] for row in rows[1:]]
        # By total, not by name: BTF-TEST-1 (14.5) before A Lower Team (2);
        # the unmarked group last.
        self.assertEqual(codes, ["BTF-TEST-1", "A Lower Team", "An Unmarked Team"])
        self.assertEqual(totals[:2], [14.5, 2])

    def test_a_supervisors_real_spreadsheet_can_be_checked_before_sending(self):
        r = self.client.get(reverse("management:results-email-supervisor-sheet", args=[self.supervisor.id]))
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        self.assertEqual(r["Content-Disposition"], f'attachment; filename="{self.group.year}_BTF_Student_Marks_Sam_Lee.xlsx"')
        rows = [[c.value for c in row] for row in load_workbook(io.BytesIO(r.content)).active.iter_rows()]
        row = dict(zip(rows[0], rows[1]))
        self.assertEqual((row["TeamCode"], row["Students"], row["MTotal"]), ("BTF-TEST-1", "Stu amy, Stu ben", 14.5))
        # The same as the one the email attaches, and nothing was sent.
        self.assertEqual(mail.outbox, [])
        self._send_all("supervisors")
        attached = _files(mail.outbox[0])[f"{self.group.year}_BTF_Student_Marks_Sam_Lee.xlsx"]
        attached_rows = [[c.value for c in row] for row in load_workbook(io.BytesIO(attached)).active.iter_rows()]
        self.assertEqual(attached_rows, rows)

    def test_only_a_supervisor_due_the_email_has_a_spreadsheet(self):
        r = self.client.get(reverse("management:results-email-supervisor-sheet", args=[self.staff.id]))
        self.assertEqual(r.status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(r.json()["detail"], "That supervisor isn't due the results email.")
        self.client.force_authenticate(self.non_staff)
        r = self.client.get(reverse("management:results-email-supervisor-sheet", args=[self.supervisor.id]))
        self.assertEqual(r.status_code, status.HTTP_403_FORBIDDEN)

    def test_the_sample_spreadsheet_is_for_graders_only(self):
        self.client.force_authenticate(self.non_staff)
        r = self.client.get(reverse("management:results-email-sample-sheet"))
        self.assertEqual(r.status_code, status.HTTP_403_FORBIDDEN)

    def test_preview_names_example_files_while_nobody_is_due(self):
        Submission.objects.update(submitted_at=None)
        r = self.client.post(reverse("management:results-email-preview"), {"audience": "groups"}, format="json")
        year = self.group.year
        self.assertEqual(r.json()["attachments"], [
            f"{year}_BTF_Student_Certificate_Name.docx",
            f"{year}_BTF_Mentor_Certificate_Name.docx",
            f"{year}_BTF_Marks_Team_name.docx",
        ])
        r = self.client.post(reverse("management:results-email-preview"), {"audience": "supervisors"}, format="json")
        self.assertEqual(r.json()["attachments"][-1], f"{year}_BTF_Student_Marks_Supervisor_name.xlsx")

    # -- Document Setup's tests with a real person ----------------------------------------

    def _people(self, kind):
        r = self.client.get(reverse("management:settings-test-people", args=[kind]))
        self.assertEqual(r.status_code, status.HTTP_200_OK, r.content)
        return r.json()["options"]

    def _value(self, kind, label):
        return next(o["value"] for o in self._people(kind) if o["label"] == label)

    def test_templates_can_be_tested_with_this_years_students_or_mentors(self):
        # Every group that submitted, finalists too; not the one with no entry.
        self.assertEqual(
            [o["label"] for o in self._people("certificate")],
            ["(BTF-TEST-1) Stu amy", "(BTF-TEST-1) Stu ben", "(Finalist Team) Stu fin"],
        )
        self.assertEqual(self._people("marks-summary"), self._people("certificate"))
        self.assertEqual([o["label"] for o in self._people("mentor-certificate")], ["(BTF-TEST-1) Mo Mentor"])
        self.assertEqual(
            self.client.get(reverse("management:settings-test-people", args=["nope"])).status_code,
            status.HTTP_404_NOT_FOUND,
        )

    def test_a_test_render_for_a_person_is_their_real_document(self):
        from apps.submissions.services import current_cohort

        year = current_cohort()
        url = reverse("management:settings-test-render", args=["certificate"])
        r = self.client.get(url, {"person": self._value("certificate", "(BTF-TEST-1) Stu amy")})
        self.assertEqual(r.status_code, status.HTTP_200_OK, r.content)
        self.assertIn(f"{year}_BTF_Student_Certificate_Stu_amy.docx", r["Content-Disposition"])
        self.assertIn("Stu amy — BTF-TEST-1", _docx_text(r.content))

        url = reverse("management:settings-test-render", args=["mentor-certificate"])
        r = self.client.get(url, {"person": self._value("mentor-certificate", "(BTF-TEST-1) Mo Mentor")})
        self.assertIn(f"{year}_BTF_Mentor_Certificate_Mo_Mentor.docx", r["Content-Disposition"])
        self.assertIn("Mo Mentor mentored BTF-TEST-1", _docx_text(r.content))

        # The marks summary is the student's group's, with its real marks.
        url = reverse("management:settings-test-render", args=["marks-summary"])
        r = self.client.get(url, {"person": self._value("marks-summary", "(BTF-TEST-1) Stu ben")})
        self.assertIn(f"{year}_BTF_Marks_BTF-TEST-1.docx", r["Content-Disposition"])
        self.assertIn("Team BTF-TEST-1 S1 8", _docx_text(r.content))

    def test_a_picked_file_is_tested_with_a_person_before_it_is_saved(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        from tests.apps.management.fixtures import _build_docx

        url = reverse("management:settings-test-render", args=["certificate"])
        r = self.client.post(url, {
            "file": SimpleUploadedFile("new.docx", _build_docx("Candidate for {{Name}}")),
            "person": self._value("certificate", "(Finalist Team) Stu fin"),
        }, format="multipart")
        self.assertEqual(r.status_code, status.HTTP_200_OK, r.content)
        self.assertIn("Candidate for Stu fin", _docx_text(r.content))
        # The saved template is untouched.
        with GradingSettings.load().certificate_template.open("rb") as fh:
            self.assertIn("{{Name}} — {{ProjectTitle}}", _docx_text(fh.read()))

    def test_a_person_not_on_the_list_is_refused(self):
        url = reverse("management:settings-test-render", args=["mentor-certificate"])
        # A student isn't on the mentor list.
        student = self._value("certificate", "(BTF-TEST-1) Stu amy")
        r = self.client.get(url, {"person": student})
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(r.json()["detail"], "That person isn't on this year's list.")
        r = self.client.get(url, {"person": "rubbish"})
        self.assertEqual(r.json()["detail"], "Pick someone from the list.")

    def test_the_people_list_is_for_graders_only(self):
        self.client.force_authenticate(_user("kid@example.com"))
        r = self.client.get(reverse("management:settings-test-people", args=["certificate"]))
        self.assertEqual(r.status_code, status.HTTP_403_FORBIDDEN)
