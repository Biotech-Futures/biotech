"""Results on a supervisor's portal: once marks or certificates are released,
a supervisor sees their students' groups' marks, as the spreadsheet their
results email carries, and the certificates of their students and of those
groups' mentors; a finalist group's certificates are released separately
while certificates go out without finalists. Only for supervisors."""
import io
import zipfile

from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from apps.grading.models import FinalistFlag
from apps.groups.models import GroupMembership, Groups
from apps.management.models import CertificatesRelease, MarksRelease
from apps.submissions.models import Submission
from apps.users.models import StudentProfile, SupervisorProfile, User

from tests.apps.grading.fixtures import _GradingFixture
from tests.apps.management.fixtures import _seed_doc_templates


def _user(email, first, last):
    return User.objects.create_user(email=email, first_name=first, last_name=last, password="pw12345!")


def _student(email, first, last, group, supervisor_profile):
    user = _user(email, first, last)
    StudentProfile.objects.create(
        user=user, pg_first_name="P", pg_last_name="G", parent_guardian_flag=True,
        school_name="Test School", year_lvl="10", supervisor=supervisor_profile,
    )
    GroupMembership.objects.create(user=user, group=group, membership_role="student")
    return user


def _submitted_team(name, staff):
    team = Groups.objects.create(group_name=name)
    submission = Submission.objects.create(group=team, answers={"q_answers": "Answers."})
    submission.snapshot(staff)
    submission.save()
    return team


def _release(model, **fields):
    row = model.load()
    row.released_at = timezone.now()
    for name, value in fields.items():
        setattr(row, name, value)
    row.save()


class SupervisorResultsTests(_GradingFixture):
    def setUp(self):
        _seed_doc_templates()
        self.sam = _user("sam.lee@example.com", "Sam", "Lee")
        sam_profile = SupervisorProfile.objects.create(user=self.sam)

        # Sam's students in the fixture's submitted group, with its mentor.
        self.amy = _student("amy@example.com", "Amy", "Chen", self.group, sam_profile)
        self.ben = _student("ben@example.com", "Ben", "Ng", self.group, sam_profile)
        self.mo = _user("mo@example.com", "Mo", "Mentor")
        GroupMembership.objects.create(user=self.mo, group=self.group, membership_role="mentor")

        # A finalist team of Sam's, and a team of Sam's that never submitted.
        self.finalist = _submitted_team("Finalist Team", self.staff)
        self.fin = _student("fin@example.com", "Fin", "Al", self.finalist, sam_profile)
        FinalistFlag.objects.create(group=self.finalist, flagged_by=self.staff)
        _student("none@example.com", "No", "Entry", Groups.objects.create(group_name="No Entry"), sam_profile)

        # Another supervisor's student, on another team with its own mentor.
        other_profile = SupervisorProfile.objects.create(user=_user("kim@example.com", "Kim", "Other"))
        self.other_team = _submitted_team("Other Team", self.staff)
        self.olly = _student("olly@example.com", "Olly", "Other", self.other_team, other_profile)
        self.other_mentor = _user("om@example.com", "Oz", "Mentor")
        GroupMembership.objects.create(user=self.other_mentor, group=self.other_team, membership_role="mentor")

    def _client(self, user=None):
        client = APIClient()
        client.force_authenticate(user or self.sam)
        return client

    def _get(self, name, *args, user=None):
        return self._client(user).get(reverse(f"management:{name}", args=args))

    def _results(self):
        response = self._get("supervisor-results")
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.content)
        return response.json()

    def test_only_supervisors_see_it(self):
        for user in (self.staff, self.amy):
            for name in ("supervisor-results", "supervisor-results-state", "supervisor-results-sheet"):
                with self.subTest(user=user.email, name=name):
                    self.assertEqual(self._get(name, user=user).status_code, status.HTTP_403_FORBIDDEN)

    def test_nothing_shows_before_anything_is_released(self):
        self.assertEqual(
            self._get("supervisor-results-state").json(),
            {"marks_released": False, "certificates_released": False},
        )
        body = self._results()
        self.assertEqual((body["groups"], body["people"], body["columns"]), ([], [], []))
        self.assertEqual(self._get("supervisor-results-sheet").status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(self._get("supervisor-results-certificates").status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(
            self._get("supervisor-results-certificate", self.group.id, self.amy.id).status_code,
            status.HTTP_403_FORBIDDEN,
        )

    def test_released_marks_show_their_students_groups_as_the_spreadsheet(self):
        _release(MarksRelease)
        self.assertEqual(self._get("supervisor-results-state").json()["marks_released"], True)

        body = self._results()
        # Their groups that submitted, finalists too; not another supervisor's.
        self.assertEqual(
            sorted(row["TeamCode"] for row in body["groups"]),
            sorted([self.group.group_name, "Finalist Team"]),
        )
        # Every column the spreadsheet has, empty or not.
        from apps.management.services.results_notify import MARKS_SHEET_COLUMNS

        self.assertEqual([c["key"] for c in body["columns"]], MARKS_SHEET_COLUMNS)
        # Certificates aren't out yet.
        self.assertEqual(body["people"], [])

        sheet = self._get("supervisor-results-sheet")
        self.assertEqual(sheet.status_code, status.HTTP_200_OK)
        self.assertIn('filename="', sheet["Content-Disposition"])
        rows = list(load_workbook(io.BytesIO(sheet.content)).active.iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "Group")
        self.assertEqual(sorted(row[0] for row in rows[1:]), sorted([self.group.group_name, "Finalist Team"]))

    def test_released_certificates_list_their_students_and_those_groups_mentors(self):
        _release(CertificatesRelease)

        body = self._results()
        people = {(p["name"], p["kind"], p["group_name"], p["withheld"]) for p in body["people"]}
        self.assertEqual(people, {
            ("Amy Chen", "student", self.group.group_name, False),
            ("Ben Ng", "student", self.group.group_name, False),
            ("Fin Al", "student", "Finalist Team", True),
            ("Mo Mentor", "mentor", self.group.group_name, False),
        })
        # Marks aren't out yet.
        self.assertEqual(body["groups"], [])

    def test_certificates_download_one_at_a_time_or_all_together(self):
        _release(CertificatesRelease)

        for holder in (self.amy, self.mo):
            with self.subTest(holder.email):
                response = self._get("supervisor-results-certificate", self.group.id, holder.id)
                self.assertEqual(response.status_code, status.HTTP_200_OK)
                self.assertTrue(zipfile.is_zipfile(io.BytesIO(response.content)))
        # A finalist's waits; anyone else's isn't theirs to see.
        self.assertEqual(
            self._get("supervisor-results-certificate", self.finalist.id, self.fin.id).status_code,
            status.HTTP_403_FORBIDDEN,
        )
        for group, holder in ((self.other_team, self.olly), (self.other_team, self.other_mentor)):
            with self.subTest(holder.email):
                self.assertEqual(
                    self._get("supervisor-results-certificate", group.id, holder.id).status_code,
                    status.HTTP_404_NOT_FOUND,
                )

        bundle = self._get("supervisor-results-certificates")
        self.assertEqual(bundle.status_code, status.HTTP_200_OK)
        names = sorted(zipfile.ZipFile(io.BytesIO(bundle.content)).namelist())
        year = self.year
        self.assertEqual(names, sorted([
            f"{year}_BTF_Student_Certificate_Amy_Chen.docx",
            f"{year}_BTF_Student_Certificate_Ben_Ng.docx",
            f"{year}_BTF_Mentor_Certificate_Mo_Mentor.docx",
        ]))

    def test_finalists_certificates_come_with_the_rest_when_not_held_back(self):
        _release(CertificatesRelease, exclude_finalists=False)

        fin = next(p for p in self._results()["people"] if p["name"] == "Fin Al")
        self.assertFalse(fin["withheld"])
        self.assertEqual(
            self._get("supervisor-results-certificate", self.finalist.id, self.fin.id).status_code,
            status.HTTP_200_OK,
        )
