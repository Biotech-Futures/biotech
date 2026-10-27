"""The Audit tab as an Excel file (client item C-26, 2026-10-22).

Two halves. The wording: the file says what the table says, and the table
words its rows in the portal, so the cases below are the ones
TicketAuditPage.spec.ts pins, with the export's plain hyphen where the table
prints a dash. And the endpoint: built as the queue's export is, so it is
held to the same things (who may, how it is asked for, what it records).
"""

import io
from datetime import datetime, timedelta, timezone as dt_timezone
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.test import SimpleTestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from apps.audit.models import AuditLog
from apps.tickets.models import SupportScope, TicketCategory
from apps.tickets.services import lifecycle
from apps.tickets.services.audit_export import action_label, actor_name, summarise

User = get_user_model()

AUDIT = "/api/v1/admin/tickets/audit/"
AUDIT_EXPORT = "/api/v1/admin/tickets/audit/export/"


class TheFileWordsRowsAsTheTableDoesTests(SimpleTestCase):
    """summarise(), action_label() and actor_name() against the table's cases."""

    names = {4: "Sam Reid"}

    def test_the_rules_run_in_the_order_delete_create_status_owner_priority_category(self):
        """A cascade, as the table's test builds it: row n carries a key for
        its own rule and every rule after it, so swapping any two rules that
        can both match one row changes what some row says."""
        every_key_before = {
            "ticket_number": "SUP-2026-00042",
            "subject": "Cannot access group workspace",
            "status": "resolved",
            "assignee_id": 4,
            "priority": "low",
            "category": "technical_issue",
        }
        every_key_after = {
            "screening_category": "self_harm",
            "status": "open",
            "assignee_id": None,
            "priority": "high",
            "category": "registration",
        }
        rows = [
            ("delete", every_key_before, every_key_after),
            ("create", every_key_before, every_key_after),
            ("status",
             {"status": "resolved", "assignee_id": 4, "priority": "low", "category": "technical_issue"},
             {"status": "open", "assignee_id": None, "priority": "high", "category": "registration"}),
            ("assign",
             {"assignee_id": None, "priority": "low", "category": "technical_issue"},
             {"assignee_id": 4, "priority": "high", "category": "registration"}),
            ("priority",
             {"priority": "low", "category": "technical_issue"},
             {"priority": "high", "category": "registration"}),
            ("category", {"category": "technical_issue"}, {"category": "registration"}),
        ]
        self.assertEqual(
            [summarise(action, before, after, self.names) for action, before, after in rows],
            [
                "SUP-2026-00042 · Cannot access group workspace",
                "Flagged as Self harm",
                "Resolved → Open",
                "Owner set to Sam Reid",
                "Low → High",
                "Technical issue → Registration",
            ],
        )

    def test_a_real_reopen_reads_as_the_status_change_and_a_real_delete_as_the_deletion(self):
        reopen = summarise(
            "reopen",
            {"status": "resolved", "assignee_id": 4, "resolved_at": "2026-09-01T09:00:00Z"},
            {"status": "open", "assignee_id": None, "resolved_at": None},
            self.names,
        )
        delete = summarise(
            "delete",
            {"ticket_number": "SUP-2026-00042", "subject": "Cannot access group workspace",
             "status": "open", "created_by_id": 7},
            {"deleted_at": "2026-09-02T09:00:00Z"},
            self.names,
        )
        self.assertEqual(
            [reopen, delete],
            ["Resolved → Open", "SUP-2026-00042 · Cannot access group workspace"],
        )

    def test_a_re_filed_ticket_says_what_it_moved_from_and_to(self):
        rows = [
            ({"category": "technical_issue"}, {"category": "registration"}),
            ({"category": "flagged_content"}, {"category": "help_student_group"}),
            ({"category": ""}, {"category": "general_question"}),
            ({"category": "other"}, {}),
        ]
        self.assertEqual(
            [summarise("category", before, after, self.names) for before, after in rows],
            [
                "Technical issue → Registration",
                "Flagged content → Help with a student or group",
                "Not categorised → General Question",
                "Other → Not categorised",
            ],
        )

    def test_a_missing_side_is_a_hyphen_and_an_unknown_status_is_as_stored(self):
        self.assertEqual(
            [
                summarise("status", None, {"status": "open"}, self.names),
                summarise("status", {"status": "resolved"}, {}, self.names),
                summarise("status", {"status": "escalated"}, {"status": "open"}, self.names),
            ],
            ["- → Open", "Resolved → -", "escalated → Open"],
        )

    def test_an_opening_without_a_verdict_and_a_row_nothing_recognises(self):
        """A priority on the after side alone is not a change: screening puts
        one on every Created row, and the priority rule reads the before side."""
        self.assertEqual(
            [
                summarise("create", None, {"channel": "portal"}, self.names),
                summarise("escalate", {}, {}, self.names),
                summarise("priority", None, {"priority": "high"}, self.names),
            ],
            ["Ticket opened", "-", "-"],
        )

    def test_owners(self):
        self.assertEqual(
            [
                summarise("assign", {"assignee_id": None}, {"assignee_id": 4}, self.names),
                summarise("assign", {"assignee_id": 4}, {"assignee_id": None}, self.names),
                # Named in the table only while on its roster; here, while the
                # account exists. Gone, the id is what is left.
                summarise("assign", {"assignee_id": None}, {"assignee_id": 99}, self.names),
            ],
            ["Owner set to Sam Reid", "Handed back to the pool", "Owner set to #99"],
        )

    def test_a_deletion_with_half_a_snapshot_or_none(self):
        self.assertEqual(
            [
                summarise("delete", {"subject": "Cannot sign in"}, {}, self.names),
                summarise("delete", None, None, self.names),
            ],
            ["Cannot sign in", "Ticket removed from the queue"],
        )

    def test_every_action_in_words_and_an_unknown_one_as_stored(self):
        actions = ["create", "assign", "status", "priority", "category", "resolve",
                   "reopen", "screening_approve", "screening_delete", "delete", "escalate"]
        self.assertEqual(
            [action_label(action) for action in actions],
            ["Created", "Assigned", "Status changed", "Priority changed", "Category changed",
             "Resolved", "Reopened", "Screening approved", "Screening message deleted",
             "Deleted", "escalate"],
        )

    def test_who_a_row_names_when_it_names_nobody(self):
        self.assertEqual(
            [
                actor_name("Grace Okafor", {}),
                actor_name(None, {"channel": "ai_screening"}),
                actor_name(None, {"channel": "portal"}),
                actor_name(None, None),
            ],
            ["Grace Okafor", "Automated screening", "Account removed", "Account removed"],
        )


class AuditExportTestCase(APITestCase):
    def setUp(self):
        cache.clear()
        self.requester = User.objects.create_user(
            email="mia@example.com", password="pass1234", first_name="Mia", last_name="Thompson",
        )
        self.agent = User.objects.create_user(
            email="agent@example.com", password="pass1234", first_name="Sam", last_name="Reid",
            timezone="Australia/Sydney",
        )
        SupportScope.objects.create(user=self.agent)
        self.other_agent = User.objects.create_user(
            email="grace@example.com", password="pass1234", first_name="Grace", last_name="Okafor",
        )
        SupportScope.objects.create(user=self.other_agent)
        self.ticket = lifecycle.create_ticket(
            user=self.requester,
            category=TicketCategory.ACCOUNT_ACCESS,
            subject="Cannot sign in",
            body="The code never arrives.",
        )
        self.client.force_login(self.agent)

    def row(self, action, *, actor=None, before=None, after=None, at=None, ticket_id=None):
        return AuditLog.objects.create(
            actor_user=actor,
            entity_type=lifecycle.AUDIT_ENTITY_TYPE,
            entity_id=ticket_id or self.ticket.pk,
            action=action,
            before_state=before,
            after_state=after,
            created_at=at or timezone.now(),
        )

    def export(self, body=None):
        return self.client.post(AUDIT_EXPORT, body or {}, format="json")

    def sheet(self, response):
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.content[:300])
        return load_workbook(io.BytesIO(response.content)).active

    def export_rows(self):
        return AuditLog.objects.filter(entity_type="ticket_audit_export")


class AuditExportContentTests(AuditExportTestCase):
    def test_the_file_is_the_table_every_row_newest_first(self):
        AuditLog.objects.filter(entity_type=lifecycle.AUDIT_ENTITY_TYPE).delete()
        base = timezone.now() - timedelta(days=1)
        self.row("assign", actor=self.agent, before={"assignee_id": None},
                 after={"assignee_id": self.other_agent.pk}, at=base)
        self.row("status", actor=self.other_agent, before={"status": "open"},
                 after={"status": "in_progress"}, at=base + timedelta(hours=1))

        ws = self.sheet(self.export())

        self.assertEqual([cell.value for cell in ws[1]], ["When", "Ticket", "Action", "Who", "What changed"])
        body = [[cell.value for cell in row[1:]] for row in ws.iter_rows(min_row=2)]
        self.assertEqual(body, [
            [f"#{self.ticket.pk}", "Status changed", "Grace Okafor", "Open → In progress"],
            [f"#{self.ticket.pk}", "Assigned", "Sam Reid", "Owner set to Grace Okafor"],
        ])

    def test_every_matching_row_not_one_page(self):
        AuditLog.objects.filter(entity_type=lifecycle.AUDIT_ENTITY_TYPE).delete()
        for minute in range(30):
            self.row("status", actor=self.agent, before={"status": "open"}, after={"status": "in_progress"},
                     at=timezone.now() - timedelta(minutes=minute))

        ws = self.sheet(self.export())

        self.assertEqual(ws.max_row - 1, 30)

    def test_times_are_the_exporters_wall_clock_with_the_zone_in_the_cell(self):
        AuditLog.objects.filter(entity_type=lifecycle.AUDIT_ENTITY_TYPE).delete()
        # 2026-10-22 03:59 UTC is 2:59 pm in Sydney (AEDT, UTC+11).
        self.row("status", actor=self.agent, before={"status": "open"}, after={"status": "in_progress"},
                 at=datetime(2026, 10, 22, 3, 59, tzinfo=dt_timezone.utc))

        when = self.sheet(self.export()).cell(row=2, column=1)

        self.assertEqual(when.value, datetime(2026, 10, 22, 14, 59))
        self.assertIn('"AEDT"', when.number_format)

    def test_the_zone_the_page_sends_is_the_one_the_file_uses(self):
        # C-37 to C-40, DEC-068: the page sends the zone its Audit table is
        # shown in, which for a profile still on UTC is the computer's.
        User.objects.filter(pk=self.agent.pk).update(timezone="UTC")
        AuditLog.objects.filter(entity_type=lifecycle.AUDIT_ENTITY_TYPE).delete()
        # 2026-10-22 03:59 UTC is 11:59 am in Perth (AWST, UTC+8).
        self.row("status", actor=self.agent, before={"status": "open"}, after={"status": "in_progress"},
                 at=datetime(2026, 10, 22, 3, 59, tzinfo=dt_timezone.utc))
        when = self.sheet(self.export({"timeZone": "Australia/Perth"})).cell(row=2, column=1)
        self.assertEqual(when.value, datetime(2026, 10, 22, 11, 59))
        self.assertIn('"AWST"', when.number_format)

    def test_an_owner_is_named_while_the_account_exists(self):
        AuditLog.objects.filter(entity_type=lifecycle.AUDIT_ENTITY_TYPE).delete()
        self.row("assign", actor=self.agent, before={"assignee_id": None}, after={"assignee_id": 987654})

        ws = self.sheet(self.export())

        self.assertEqual(ws.cell(row=2, column=5).value, "Owner set to #987654")

    def test_a_typed_formula_stays_text(self):
        """A deleted ticket's subject is a student's words, written into the
        snapshot the deletion kept."""
        AuditLog.objects.filter(entity_type=lifecycle.AUDIT_ENTITY_TYPE).delete()
        self.row("delete", actor=self.agent,
                 before={"subject": '=HYPERLINK("http://example.com","Open")'}, after={})

        cell = self.sheet(self.export()).cell(row=2, column=5)

        self.assertEqual(cell.value, '=HYPERLINK("http://example.com","Open")')
        self.assertEqual(cell.data_type, "s")
        self.assertTrue(cell.quotePrefix)

    def test_the_file_is_named_for_the_exporters_date(self):
        """2:00 pm UTC on the 22nd is 1:00 am on the 23rd in Sydney: a Sydney
        agent exporting then gets a file named for their own day."""
        with patch("django.utils.timezone.now",
                   return_value=datetime(2026, 10, 22, 14, 0, tzinfo=dt_timezone.utc)):
            response = self.export()

        self.assertEqual(response["Content-Disposition"],
                         'attachment; filename="ticket-audit-2026-10-23.xlsx"')
        self.assertTrue(response["Content-Type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))


class AuditExportFilterTests(AuditExportTestCase):
    def setUp(self):
        super().setUp()
        AuditLog.objects.filter(entity_type=lifecycle.AUDIT_ENTITY_TYPE).delete()
        self.row("assign", actor=self.agent, before={"assignee_id": None}, after={"assignee_id": self.agent.pk})
        self.row("status", actor=self.agent, before={"status": "open"}, after={"status": "in_progress"})
        self.row("status", actor=self.other_agent, before={"status": "in_progress"}, after={"status": "resolved"})

    def actions_in(self, body):
        ws = self.sheet(self.export(body))
        return [(row[2].value, row[3].value) for row in ws.iter_rows(min_row=2)]

    def test_the_tabs_action_filter(self):
        self.assertEqual(
            sorted(self.actions_in({"action": "status"})),
            [("Status changed", "Grace Okafor"), ("Status changed", "Sam Reid")],
        )

    def test_the_tabs_who_filter(self):
        self.assertEqual(self.actions_in({"actor": str(self.other_agent.pk)}),
                         [("Status changed", "Grace Okafor")])

    def test_the_file_and_the_table_agree_on_what_matches(self):
        filters = {"action": "status", "actor": str(self.agent.pk)}
        table = self.client.get(AUDIT, filters).json()["data"]["total"]

        self.assertEqual(len(self.actions_in(filters)), table)

    def test_a_filter_that_is_not_text_is_refused(self):
        for body in ({"action": ["status"]}, {"actor": self.agent.pk}):
            with self.subTest(body=body):
                response = self.export(body)

                self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
                self.assertFalse(self.export_rows().exists())

    def test_a_body_that_is_not_an_object_is_refused(self):
        response = self.client.post(AUDIT_EXPORT, ["status"], format="json")

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_a_huge_actor_id_is_a_400_not_a_500(self):
        response = self.export({"actor": str(2 ** 64)})

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)


class AuditExportAccessTests(AuditExportTestCase):
    def test_somebody_who_cannot_read_the_tab_cannot_take_it_away(self):
        self.client.force_login(self.requester)

        response = self.export()

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertFalse(self.export_rows().exists())

    def test_signed_out_is_refused(self):
        self.client.logout()

        self.assertEqual(self.export().status_code, status.HTTP_403_FORBIDDEN)

    def test_a_get_is_refused_and_records_nothing(self):
        """POST only, for the queue export's reason: a GET is not CSRF-checked."""
        response = self.client.get(AUDIT_EXPORT)

        self.assertEqual(response.status_code, status.HTTP_405_METHOD_NOT_ALLOWED)
        self.assertFalse(self.export_rows().exists())


class AuditExportRecordTests(AuditExportTestCase):
    def test_the_export_is_recorded_with_its_filters_and_size(self):
        AuditLog.objects.filter(entity_type=lifecycle.AUDIT_ENTITY_TYPE).delete()
        self.row("status", actor=self.agent, before={"status": "open"}, after={"status": "in_progress"})

        self.export({"action": "status", "actor": ""})

        record = self.export_rows().get()
        self.assertEqual(record.actor_user_id, self.agent.pk)
        self.assertEqual(record.entity_id, self.agent.pk)
        self.assertEqual(record.action, AuditLog.ActionChoices.EXPORT)
        self.assertEqual(record.after_state, {"filters": {"action": "status"}, "rowCount": 1})

    def test_the_record_is_not_listed_on_the_tab_itself(self):
        self.export()

        listed = {row["action"] for row in self.client.get(AUDIT).json()["data"]["items"]}
        self.assertNotIn("export", listed)
