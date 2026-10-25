"""Upload students by CSV, from an .xlsx: a supervisor's spreadsheet comes
back as CSV text for the page to check as it checks a CSV."""
import io

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook
from rest_framework.test import APIClient

from tests.apps.users.test_registration_multi_student import MultiStudentSetup

URL = "/api/v1/registration/sheet/"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx(rows) -> SimpleUploadedFile:
    book = Workbook()
    for row in rows:
        book.active.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return SimpleUploadedFile("students.xlsx", buffer.getvalue(), content_type=XLSX)


class RegistrationSheetTests(MultiStudentSetup, TestCase):
    def test_a_supervisors_xlsx_comes_back_as_csv(self):
        upload = _xlsx([
            ["first_name", "last_name", "email", "year_level", "interests"],
            ["Amy", "Chen", "amy@example.test", 10, "Water, Energy"],
            ["Ben", None, "ben@example.test", 11.0, None],
        ])
        response = self.supervisor_client.post(URL, {"file": upload}, format="multipart")
        self.assertEqual(response.status_code, 200, response.data)
        # Whole numbers without ".0", empty cells empty, commas quoted.
        self.assertEqual(response.data["csv"], (
            "first_name,last_name,email,year_level,interests\n"
            'Amy,Chen,amy@example.test,10,"Water, Energy"\n'
            "Ben,,ben@example.test,11,\n"
        ))

    def test_only_a_supervisor_and_only_a_readable_file(self):
        self.assertEqual(APIClient().post(URL, {"file": _xlsx([["a"]])}, format="multipart").status_code, 403)

        broken = SimpleUploadedFile("students.xlsx", b"not a workbook", content_type=XLSX)
        response = self.supervisor_client.post(URL, {"file": broken}, format="multipart")
        self.assertEqual(response.status_code, 400)
        self.assertIn("file", response.data["fields"])

        response = self.supervisor_client.post(URL, {}, format="multipart")
        self.assertEqual(response.status_code, 400)

        big = SimpleUploadedFile("students.xlsx", b"0" * (2 * 1024 * 1024 + 1), content_type=XLSX)
        response = self.supervisor_client.post(URL, {"file": big}, format="multipart")
        self.assertEqual(response.status_code, 400)
        self.assertIn("2 MB", str(response.data["fields"]["file"]))


TEMPLATE = "/api/v1/registration/sheet/template/"


class RegistrationTemplateTests(MultiStudentSetup, TestCase):
    def template(self, **changes):
        return {
            "header": ["First Name", "Email"], "example": ["Alex", "alex.example@school.invalid"],
            "note": "Example only: this row won't be registered.", **changes,
        }

    def test_the_template_is_an_xlsx_with_its_example_row_and_note(self):
        from openpyxl import load_workbook

        response = self.supervisor_client.post(TEMPLATE, self.template(), format="json")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], XLSX)
        self.assertIn('filename="biotech-registration-template.xlsx"', response["Content-Disposition"])
        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual([[cell.value for cell in row] for row in sheet.iter_rows()], [
            ["First Name", "Email", None],
            ["Alex", "alex.example@school.invalid", "Example only: this row won't be registered."],
        ])
        self.assertTrue(sheet["A1"].font.bold)
        # Read back as an upload is, the note is only an extra column.
        upload = SimpleUploadedFile("students.xlsx", response.content, content_type=XLSX)
        read = self.supervisor_client.post(URL, {"file": upload}, format="multipart")
        self.assertEqual(read.data["csv"].splitlines()[1], "Alex,alex.example@school.invalid,Example only: this row won't be registered.")

    def test_only_a_supervisor_and_a_value_for_each_column(self):
        self.assertEqual(APIClient().post(TEMPLATE, self.template(), format="json").status_code, 403)
        response = self.supervisor_client.post(TEMPLATE, self.template(example=["Alex"]), format="json")
        self.assertEqual(response.status_code, 400)
